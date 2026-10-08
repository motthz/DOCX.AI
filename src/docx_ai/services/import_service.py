"""Import di interventi in blocco da CSV o Excel.

1. ``read_table`` legge intestazioni e righe (CSV con separatore rilevato
   automaticamente, oppure il primo foglio di un .xlsx).
2. ``suggest_mapping`` associa le colonne ai campi dello schema del modulo
   confrontando i nomi normalizzati.
3. ``import_rows`` crea un rapporto per riga (bozza o approvato), convertendo
   i valori secondo il tipo del campo.
"""

from __future__ import annotations

import csv
import difflib
import json
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from ..db import Database
from ..module_manager import LoadedModule

TRUE_WORDS = {"si", "sì", "s", "x", "true", "vero", "1", "yes", "y", "ok"}
FALSE_WORDS = {"no", "n", "false", "falso", "0", ""}


@dataclass
class Table:
    headers: List[str]
    rows: List[List[Any]] = field(default_factory=list)


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def read_table(path: Path) -> Table:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            values = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        raw = path.read_bytes()
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        text = text or ""
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";" if text.count(";") > text.count(",") else ","
        values = [row for row in csv.reader(text.splitlines(), dialect)]
    values = [v for v in values if any(c not in (None, "") for c in v)]
    if not values:
        return Table([])
    headers = [str(h or f"colonna_{i + 1}").strip() for i, h in enumerate(values[0])]
    return Table(headers, values[1:])


def suggest_mapping(headers: List[str], schema: Dict[str, Any]) -> Dict[str, Optional[str]]:
    """{campo_schema: intestazione_colonna | None}"""
    fields = list((schema or {}).get("properties", {}).keys())
    normed = {_norm(h): h for h in headers}
    out: Dict[str, Optional[str]] = {}
    used: set = set()
    for f in fields:
        nf = _norm(f)
        title = _norm((schema["properties"][f] or {}).get("title", ""))
        hit = normed.get(nf) or (normed.get(title) if title else None)
        if hit is None:
            close = difflib.get_close_matches(nf, list(normed), n=1, cutoff=0.72)
            hit = normed[close[0]] if close else None
        if hit in used:
            hit = None
        out[f] = hit
        if hit:
            used.add(hit)
    # secondo passaggio, solo abbinamenti univoci: "Giorni" -> giorni_lavorativi (parole
    # contenute nel nome del campo), "Dal"/"Al" -> data_inizio/data_fine
    free = [h for h in headers if h not in used]
    for h in free:
        hw = set(_norm(h).split("_")) - {""}
        if not hw:
            continue
        cands = []
        for f in fields:
            if out.get(f):
                continue
            spec = (schema["properties"][f] or {})
            fw = set(_norm(f).split("_")) | set(_norm(spec.get("title", "")).split("_"))
            if (all(len(w) >= 4 for w in hw) and hw <= fw) or any(
                    hw <= syn and fw & targets for syn, targets in _PERIOD_SYNONYMS):
                cands.append(f)
        if len(cands) == 1:
            out[cands[0]] = h
            used.add(h)
    return out


# intestazioni brevi di un periodo e parole dei campi a cui corrispondono
_PERIOD_SYNONYMS = (({"dal", "da", "dalla", "inizio"}, {"inizio", "dal", "partenza"}),
                    ({"al", "a", "alla", "fino", "fine"}, {"fine", "al", "termine", "rientro"}))


def _convert(value: Any, spec: Dict[str, Any]) -> Any:
    t = spec.get("type")
    if isinstance(t, list):
        t = next((x for x in t if x != "null"), "string")
    if value is None or (isinstance(value, str) and not value.strip()):
        return False if t == "boolean" else "NON_SPECIFICATO" if t in (None, "string") else None
    if t == "boolean":
        return str(value).strip().lower() in TRUE_WORDS
    if t in ("number", "integer"):
        try:
            n = float(str(value).replace(",", "."))
            return int(n) if t == "integer" else n
        except ValueError:
            return None
    if t == "array":
        if isinstance(value, str):
            return [v.strip() for v in re.split(r"[;\n]", value) if v.strip()]
        return [value]
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    if spec.get("format") == "date":
        m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})$", s)
        if m:
            d, mo, y = m.groups()
            y = int(y) + (2000 if len(y) == 2 else 0)
            return f"{y:04d}-{int(mo):02d}-{int(d):02d}"
    enum = spec.get("enum")
    if enum and s not in enum:
        match = difflib.get_close_matches(_norm(s), [_norm(e) for e in enum], n=1, cutoff=0.6)
        if match:
            return enum[[_norm(e) for e in enum].index(match[0])]
    return s


def import_rows(db: Database, mod: LoadedModule, table: Table, mapping: Dict[str, Optional[str]], *,
                status: str = "draft", description_column: Optional[str] = None) -> List[int]:
    props = (mod.schema or {}).get("properties", {})
    idx = {h: i for i, h in enumerate(table.headers)}
    ids: List[int] = []
    for row in table.rows:
        data: Dict[str, Any] = {}
        for fname, spec in props.items():
            col = mapping.get(fname)
            value = row[idx[col]] if col in idx and idx[col] < len(row) else None
            data[fname] = _convert(value, spec or {})
        desc = ""
        if description_column and description_column in idx and idx[description_column] < len(row):
            desc = str(row[idx[description_column]] or "")
        payload = json.dumps(data, ensure_ascii=False)
        rid = db.create_report(
            module_id=mod.id, module_version=mod.version or "1.0.0", status=status,
            input_description=desc or "Importato da file", draft_json=payload,
            final_json=payload if status == "approved" else None, source="import",
        )
        if status == "approved":
            db.add_report_version(rid, payload, "Importazione")
        ids.append(rid)
    return ids
