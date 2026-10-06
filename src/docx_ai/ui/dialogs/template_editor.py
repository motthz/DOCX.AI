"""Editor visuale dei moduli: si trascinano i campi nei punti che l'AI deve compilare.

- A sinistra le "tessere" dei tipi di campo (Testo, Numero, Data, Sì/No, Scelta,
  Elenco) e i campi gia' creati: si trascinano sul documento. In alternativa un
  clic sulla tessera e poi un clic nel documento.
- Al centro il documento: Word (testo e tabelle, i campi sono "pillole") oppure
  Excel (griglia delle celle). Rilasciando su una riga da compilare (______), su
  una cella vuota o su un testo selezionato il campo prende quel posto.
- I campi nel documento si spostano trascinandoli; nel cestino si tolgono.
- A destra le proprieta' del campo selezionato (nome, tipo, istruzioni per l'AI).
- "Annulla" ripristina l'ultima modifica; il salvataggio conserva una copia del
  template precedente in _versions.
"""

from __future__ import annotations

import copy
import io
import logging
import os
import shutil
import tempfile
import threading
import time
import tkinter as tk
from pathlib import Path
from tkinter import messagebox
from typing import Any, Dict, List, Optional, Tuple

import customtkinter as ctk

from ... import template_tools as tt
from ...parsers.docx_parser import _iter_all_paragraphs, _iter_block_items
from ...template_tools import replace_span  # noqa: F401 - compatibilita' con chi lo importava da qui
from ..design import C, col, font
from ..dnd import DragController
from ..i18n import t
from ..icons import icon
from ..widgets import Chip, button, label
from .base import Dialog

LOG = logging.getLogger(__name__)

# tipo -> (nome, icona, esempio)
KINDS: Dict[str, Tuple[str, str, str]] = {
    "text": ("Testo", "type", "Nomi, indirizzi, frasi"),
    "number": ("Numero", "hash", "Quantità, importi, ore"),
    "date": ("Data", "calendar", "Giorno, scadenza"),
    "bool": ("Sì / No", "square-check", "Casella da spuntare"),
    "choice": ("Scelta", "list-checks", "Una voce da un elenco fisso"),
    "list": ("Elenco", "list-todo", "Più righe: materiali, attività…"),
}


def open_visual_editor(win: Any, mod, *, first_time: bool = False) -> Optional[Dialog]:
    """Apre l'editor adatto al template del modulo (Word o Excel)."""
    if not mod.template_path or not Path(mod.template_path).exists():
        win.toast(t("Il modulo non ha un template: aggiungi un file DOCX o XLSX nella cartella del modulo."),
                  "warning")
        return None
    if mod.template_type == "xlsx":
        return SheetEditor(win, mod, first_time=first_time)
    from ... import docx_layout as dl
    if dl.converter().available:  # pagine vere, impaginate da Word o LibreOffice
        return PageTemplateEditor(win, mod, first_time=first_time)
    win.toast(t("Per vedere il modulo esattamente com'è serve Microsoft Word o LibreOffice (gratuito): "
                "per ora uso la vista semplificata."), "warning",
              action=(t("Scarica LibreOffice"), _open_libreoffice_download))
    return TemplateEditor(win, mod, first_time=first_time)


def _open_libreoffice_download() -> None:
    import webbrowser
    webbrowser.open("https://it.libreoffice.org/download/download/")


# ====================================================================== nuovo campo
class FieldDialog(Dialog):
    """Domande minime per un nuovo campo: come si chiama e (facoltativo) cosa scrivere."""

    def __init__(self, master, kind: str, suggestion: str = ""):
        name, ic, _ex = KINDS[kind]
        super().__init__(master, t("Nuovo campo: {k}", k=t(name)), width=540, height=500 if kind == "choice" else 420,
                         icon_name=ic, subtitle=t("Dai un nome al campo: l'AI lo userà per capire cosa scrivere."))
        self.kind = kind
        label(self.body, t("Come si chiama questo campo?"), kind="body_b").pack(anchor="w")
        self.name = ctk.CTkEntry(self.body, height=36, font=font("body"),
                                 placeholder_text=t("es. Nome del cliente, Data della visita"))
        self.name.pack(fill="x", pady=(4, 12))
        if suggestion:
            self.name.insert(0, suggestion)
        label(self.body, t("Cosa deve scrivere l'AI qui? (facoltativo)"), kind="body_b").pack(anchor="w")
        self.desc = ctk.CTkEntry(self.body, height=36,
                                 placeholder_text=t("es. Nome e cognome completi, come nel documento d'identità"))
        self.desc.pack(fill="x", pady=(4, 12))
        self.options: Optional[ctk.CTkTextbox] = None
        if kind == "choice":
            label(self.body, t("Opzioni tra cui scegliere (una per riga)"), kind="body_b").pack(anchor="w")
            self.options = ctk.CTkTextbox(self.body, height=90, border_width=1, border_color=C["border"])
            self.options.pack(fill="x", pady=(4, 12))
            self.options.insert("1.0", "Sì\nNo")
        self.required = ctk.CTkCheckBox(self.body, text=t("Obbligatorio: segnalamelo se l'AI non lo trova"))
        self.required.pack(anchor="w")
        button(self.footer, t("Aggiungi campo"), self._ok, kind="primary", icon_name="check").pack(
            side="right", padx=(8, 20), pady=12)
        button(self.footer, t("Annulla"), self.cancel).pack(side="right", pady=12)
        self.bind("<Return>", lambda e: None if isinstance(e.widget, tk.Text) else self._ok())
        self.after(350, self._focus)

    def _focus(self) -> None:
        try:
            self.name.focus_set()
            self.name.select_range(0, "end")
        except tk.TclError:
            pass

    def _ok(self) -> None:
        name = self.name.get().strip()
        if not name:
            self.name.focus_set()
            return
        opts = None
        if self.options is not None:
            opts = [o.strip() for o in self.options.get("1.0", "end").splitlines() if o.strip()]
        self.result = {"title": name, "description": self.desc.get().strip(), "options": opts,
                       "required": bool(self.required.get())}
        self.close()


# ====================================================================== base comune
class _VisualEditor(Dialog):
    CANVAS_HINT = ""

    def __init__(self, win: Any, mod, title: str, *, first_time: bool = False):
        super().__init__(win.root, title, width=1340, height=860, icon_name="layout-template",
                         subtitle=t("Trascina i campi nei punti del documento che l'AI deve compilare."))
        self.win = win
        self.mod = mod
        self.schema: Dict[str, Any] = copy.deepcopy(mod.schema or {"type": "object", "properties": {}})
        self.schema.setdefault("properties", {})
        self.schema.setdefault("required", [])
        self.dirty = False
        self.undo_stack: List[Tuple[Dict[str, Any], Any]] = []
        self.selected: Optional[str] = None
        self.armed: Optional[Dict[str, Any]] = None
        self._tiles: Dict[str, ctk.CTkFrame] = {}
        self.drag = DragController(self)

        b = self.body
        b.columnconfigure(1, weight=1)
        b.rowconfigure(1, weight=1)
        self._build_toolbar(b)
        self._build_palette(b)
        self.canvas_box = ctk.CTkFrame(b, fg_color="transparent")
        self.canvas_box.grid(row=1, column=1, sticky="nsew", padx=12)
        self.props = ctk.CTkScrollableFrame(b, width=290, fg_color=C["surface"], corner_radius=10,
                                            border_width=1, border_color=C["border"])
        self.props.grid(row=1, column=2, sticky="nsew")

        button(self.footer, t("Salva"), self.save, kind="primary", icon_name="save").pack(
            side="right", padx=(8, 20), pady=12)
        self._footer_extra()
        button(self.footer, t("Chiudi"), self.cancel, kind="ghost").pack(side="right", padx=8, pady=12)
        self.status = label(self.footer, "", kind="small", muted=True)
        self.status.pack(side="left", padx=20)
        self.bind("<Control-z>", lambda e: self.undo())
        self.first_time = first_time

    # ------------------------------------------------------------ struttura
    def _build_toolbar(self, b) -> None:
        bar = ctk.CTkFrame(b, fg_color="transparent")
        bar.grid(row=0, column=1, columnspan=2, sticky="ew", padx=(12, 0), pady=(0, 8))
        self.undo_btn = button(bar, t("Annulla modifica"), self.undo, icon_name="undo-2", kind="secondary",
                               height=30, state="disabled")
        self.undo_btn.pack(side="left")
        self.toolbar = bar
        self.banner = ctk.CTkLabel(bar, text="", font=font("small_b"), text_color=C["link"],
                                   fg_color=C["info_soft"], corner_radius=8, height=30)

    def _build_palette(self, b) -> None:
        left = ctk.CTkFrame(b, fg_color="transparent")
        left.grid(row=0, column=0, rowspan=2, sticky="nsew")
        left.rowconfigure(0, weight=1)
        left.columnconfigure(0, weight=1)
        pal = ctk.CTkScrollableFrame(left, width=240, fg_color=C["surface"], corner_radius=10, border_width=1,
                                     border_color=C["border"])
        pal.grid(row=0, column=0, sticky="nsew")
        label(pal, t("1. Trascina un campo"), kind="h4").pack(anchor="w", padx=10, pady=(10, 0))
        label(pal, t("Rilascialo nel punto del documento che l'AI deve compilare."), kind="caption",
              muted=True, wraplength=220).pack(anchor="w", padx=10, pady=(0, 6))
        for kind, (name, ic, example) in KINDS.items():
            tile = ctk.CTkFrame(pal, fg_color=C["surface_alt"], corner_radius=10, border_width=1,
                                border_color=C["border"])
            tile.pack(fill="x", padx=8, pady=3)
            ctk.CTkLabel(tile, text="", image=icon("grip-vertical", 14)).pack(side="left", padx=(6, 0))
            ctk.CTkLabel(tile, text="", image=icon(ic, 20, "primary"), width=34, height=34, corner_radius=8,
                         fg_color=C["primary_soft"]).pack(side="left", padx=6, pady=6)
            txt = ctk.CTkFrame(tile, fg_color="transparent")
            txt.pack(side="left", fill="x", expand=True)
            ctk.CTkLabel(txt, text=t(name), font=font("body_b"), text_color=C["text"], anchor="w",
                         height=18).pack(fill="x")
            ctk.CTkLabel(txt, text=t(example), font=font("caption"), text_color=C["text_muted"], anchor="w",
                         height=16).pack(fill="x", pady=(0, 2))
            payload = {"kind": "new", "type": kind}
            self.drag.add_source(tile, payload, "＋ " + t(name), on_click=lambda p=payload: self.arm(p))
            self._tiles[kind] = tile
        label(pal, t("I campi del modulo"), kind="h4").pack(anchor="w", padx=10, pady=(16, 0))
        label(pal, t("Clic per modificarli, trascinali per metterli (anche in più punti)."), kind="caption",
              muted=True, wraplength=220).pack(anchor="w", padx=10, pady=(0, 4))
        self.fields_box = ctk.CTkFrame(pal, fg_color="transparent")
        self.fields_box.pack(fill="x", padx=4)
        # il cestino resta sempre visibile sotto il pannello (non scorre con l'elenco)
        self.trash = ctk.CTkFrame(left, fg_color=C["surface_alt"], corner_radius=10, border_width=2,
                                  border_color=C["border"], height=74)
        self.trash.grid(row=1, column=0, sticky="ew", pady=(8, 0))
        self.trash.grid_propagate(False)
        ctk.CTkLabel(self.trash, text=t("  Trascina qui per togliere un campo"), image=icon("trash-2", 18),
                     compound="left", font=font("small"), text_color=C["text_muted"]).place(
            relx=0.5, rely=0.5, anchor="center")
        self.drag.add_target(self.trash, hover=self._trash_hover, drop=self._trash_drop, leave=self._trash_leave)

    def _footer_extra(self) -> None:
        pass

    # ------------------------------------------------------------ elenco dei campi
    def placements(self) -> Dict[str, int]:
        """Quante volte ogni campo compare nel documento (da implementare)."""
        return {}

    def render_fields(self) -> None:
        for w in self.fields_box.winfo_children():
            w.destroy()
        used = self.placements()
        props = self.schema.get("properties", {})
        if not props:
            label(self.fields_box, t("Nessun campo ancora: trascina una tessera qui sopra nel documento."),
                  kind="caption", muted=True, wraplength=210).pack(anchor="w", padx=8, pady=4)
        for key, spec in props.items():
            on = key == self.selected
            row = ctk.CTkFrame(self.fields_box, fg_color=C["selection"] if on else "transparent", corner_radius=8)
            row.pack(fill="x", pady=1)
            ctk.CTkLabel(row, text="", image=icon("grip-vertical", 14)).pack(side="left", padx=(4, 0))
            ic = KINDS[tt.kind_of(spec)][1]
            ctk.CTkLabel(row, text="", image=icon(ic, 16)).pack(side="left", padx=4)
            ctk.CTkLabel(row, text=_short(tt.title_of(key, spec), 20), font=font("small_b" if on else "small"),
                         text_color=C["text"], anchor="w").pack(side="left", fill="x", expand=True, pady=4)
            n = used.get(key, 0)
            Chip(row, "✓" if n else t("da mettere"), "success" if n else "warning").pack(side="right", padx=4)
            self.drag.add_source(row, {"kind": "field", "key": key}, tt.title_of(key, spec),
                                 on_click=lambda k=key: self.select(k))
        missing = sorted(k for k in used if k not in props)
        if missing:
            box = ctk.CTkFrame(self.fields_box, fg_color=C["warning_soft"], corner_radius=8)
            box.pack(fill="x", pady=(8, 2))
            label(box, t("Nel documento ci sono campi sconosciuti: {f}", f=", ".join(missing)), kind="caption",
                  wraplength=200).pack(fill="x", padx=8, pady=(6, 2))
            button(box, t("Aggiungili ai campi"), self.add_missing, height=26, icon_name="plus").pack(
                fill="x", padx=8, pady=(0, 8))
        self._update_status()

    def add_missing(self) -> None:
        props = self.schema["properties"]
        missing = [k for k in self.placements() if k not in props]
        if not missing:
            return
        self.snapshot()
        for k in missing:
            props[k] = {"type": "string"}
        self.changed()

    def _update_status(self) -> None:
        props = self.schema.get("properties", {})
        used = self.placements()
        todo = [k for k in props if not used.get(k)]
        msg = t("{n} campi", n=len(props))
        if todo:
            msg += " · " + t("{n} ancora da mettere nel documento", n=len(todo))
        if self.dirty:
            msg += " · " + t("modifiche non salvate")
        self.status.configure(text=msg)

    # ------------------------------------------------------------ pannello proprieta'
    def select(self, key: Optional[str]) -> None:
        self.selected = key if key in self.schema.get("properties", {}) else None
        self.render_fields()
        self.render_props()
        self.highlight_selected()

    def highlight_selected(self) -> None:
        pass

    def render_props(self) -> None:
        for w in self.props.winfo_children():
            w.destroy()
        p = self.props
        key = self.selected
        if key is None:
            self._render_guide(p)
            return
        spec = self.schema["properties"][key]
        kind = tt.kind_of(spec)
        label(p, t("Campo selezionato"), kind="h4").pack(anchor="w", padx=12, pady=(12, 8))
        label(p, t("Nome"), kind="small_b").pack(anchor="w", padx=12)
        name = ctk.CTkEntry(p, height=32)
        name.insert(0, tt.title_of(key, spec))
        name.pack(fill="x", padx=12, pady=(2, 10))
        label(p, t("Tipo"), kind="small_b").pack(anchor="w", padx=12)
        names = {t(v[0]): k for k, v in KINDS.items()}
        kind_menu = ctk.CTkOptionMenu(p, values=list(names), height=32)
        kind_menu.set(t(KINDS[kind][0]))
        kind_menu.pack(fill="x", padx=12, pady=(2, 10))
        label(p, t("Cosa deve scrivere l'AI?"), kind="small_b").pack(anchor="w", padx=12)
        desc = ctk.CTkTextbox(p, height=80, border_width=1, border_color=C["border"], wrap="word")
        desc.insert("1.0", str(spec.get("description") or ""))
        desc.pack(fill="x", padx=12, pady=(2, 10))
        opt_lbl = label(p, t("Opzioni (una per riga)"), kind="small_b")
        opts = ctk.CTkTextbox(p, height=80, border_width=1, border_color=C["border"])
        opts.insert("1.0", "\n".join(str(v) for v in spec.get("enum") or []))
        req = ctk.CTkSwitch(p, text=t("Obbligatorio"))
        if key in self.schema.get("required", []):
            req.select()

        def toggle_opts(_v=None) -> None:
            if names.get(kind_menu.get()) == "choice":
                opt_lbl.pack(anchor="w", padx=12, before=req)
                opts.pack(fill="x", padx=12, pady=(2, 10), before=req)
            else:
                opt_lbl.pack_forget()
                opts.pack_forget()

        kind_menu.configure(command=toggle_opts)
        req.pack(anchor="w", padx=12, pady=(0, 10))
        toggle_opts()
        label(p, self.where_text(key), kind="caption", muted=True, wraplength=260).pack(
            anchor="w", padx=12, pady=(0, 10))

        def apply() -> None:
            new_kind = names.get(kind_menu.get(), kind)
            options = [o.strip() for o in opts.get("1.0", "end").splitlines() if o.strip()]
            if new_kind == "choice" and not options:
                self.win.toast(t("Scrivi almeno un'opzione, una per riga."), "warning")
                return
            self.snapshot()
            new = tt.change_kind(spec, new_kind, options if new_kind == "choice" else None)
            title = name.get().strip()
            if title and title != tt.title_of(key, None):
                new["title"] = title
            else:
                new.pop("title", None)
            d = desc.get("1.0", "end").strip()
            if d:
                new["description"] = d
            else:
                new.pop("description", None)
            self.schema["properties"][key] = new
            required = [r for r in self.schema.get("required", []) if r != key]
            if req.get():
                required.append(key)
            self.schema["required"] = required
            self.changed()
            self.win.toast(t("Campo aggiornato."), "success")

        button(p, t("Applica modifiche"), apply, kind="primary", icon_name="check").pack(
            fill="x", padx=12, pady=(4, 6))
        if self.placements().get(key):
            button(p, t("Togli dal documento"), lambda: self.unplace_all(key), icon_name="x").pack(
                fill="x", padx=12, pady=3)
        button(p, t("Elimina campo"), lambda: self.delete_field(key), kind="ghost", icon_name="trash-2",
               text_color=C["danger"]).pack(fill="x", padx=12, pady=(3, 12))

    def _render_guide(self, p) -> None:
        label(p, t("Come funziona"), kind="h4").pack(anchor="w", padx=12, pady=(12, 8))
        steps = (
            ("grip-vertical", t("Trascina una tessera (Testo, Data…) dal pannello di sinistra.")),
            ("mouse-pointer-click", self.CANVAS_HINT),
            ("pencil", t("Scrivi il nome del campo: l'AI capirà cosa inserire.")),
            ("save", t("Premi Salva. Da ora l'AI compilerà quei punti.")),
        )
        for i, (ic, txt) in enumerate(steps, 1):
            r = ctk.CTkFrame(p, fg_color="transparent")
            r.pack(fill="x", padx=10, pady=4)
            ctk.CTkLabel(r, text=str(i), width=26, height=26, corner_radius=13, fg_color=C["primary_soft"],
                         text_color=C["link"], font=font("small_b")).pack(side="left", anchor="n")
            label(r, txt, kind="small", wraplength=210).pack(side="left", padx=8, fill="x")
        tips = ctk.CTkFrame(p, fg_color=C["surface_alt"], corner_radius=8)
        tips.pack(fill="x", padx=10, pady=(12, 10))
        label(tips, t("Consigli"), kind="small_b").pack(anchor="w", padx=10, pady=(8, 2))
        for tip in self.tips():
            label(tips, "• " + tip, kind="caption", muted=True, wraplength=230).pack(anchor="w", padx=10, pady=1)
        ctk.CTkFrame(tips, height=6, fg_color="transparent").pack()

    def tips(self) -> List[str]:
        return [t("Non sai trascinare? Clicca una tessera e poi clicca nel documento."),
                t("Hai sbagliato? Premi «Annulla modifica» (Ctrl+Z)."),
                t("Per togliere un campo trascinalo nel cestino.")]

    def where_text(self, key: str) -> str:
        return ""

    # ------------------------------------------------------------ modalita' "clic e clic"
    def arm(self, payload: Dict[str, Any]) -> None:
        if self.armed == payload:
            self.disarm()
            return
        self.armed = payload
        for k, tile in self._tiles.items():
            on = payload.get("type") == k
            tile.configure(border_color=C["primary"] if on else C["border"],
                           fg_color=C["primary_soft"] if on else C["surface_alt"])
        self.banner.configure(text="  " + t("Ora clicca nel documento dove vuoi il campo «{k}» (Esc per annullare)",
                                            k=t(KINDS[payload["type"]][0])) + "  ")
        self.banner.pack(side="left", padx=12)
        self.on_arm(True)

    def disarm(self) -> None:
        self.armed = None
        self.first_time = False
        for tile in self._tiles.values():
            tile.configure(border_color=C["border"], fg_color=C["surface_alt"])
        self.banner.pack_forget()
        self.on_arm(False)

    def on_arm(self, on: bool) -> None:
        pass

    def cancel(self) -> None:
        if self.drag.dragging:
            self.drag.cancel()
            return
        if self.armed is not None:
            self.disarm()
            return
        if self.dirty and not messagebox.askyesno(t("Modifiche non salvate"),
                                                  t("Chiudere senza salvare le modifiche al modulo?"), parent=self):
            return
        super().cancel()

    # ------------------------------------------------------------ cestino
    def _trash_hover(self, _x, _y, payload) -> bool:
        ok = payload.get("kind") in ("placed", "field")
        self.trash.configure(border_color=C["danger"] if ok else C["border"],
                             fg_color=C["danger_soft"] if ok else C["surface_alt"])
        return ok

    def _trash_leave(self) -> None:
        self.trash.configure(border_color=C["border"], fg_color=C["surface_alt"])

    def _trash_drop(self, _x, _y, payload) -> None:
        if payload.get("kind") == "placed":
            self.snapshot()
            self.unplace(payload)
            self.changed()
        elif payload.get("kind") == "field":
            self.delete_field(payload["key"])

    # ------------------------------------------------------------ campi
    def create_field(self, kind: str, suggestion: str = "") -> Optional[str]:
        dlg = FieldDialog(self, kind, suggestion)
        res = dlg.wait()
        self.after(50, self._regrab)
        if not res:
            return None
        props = self.schema["properties"]
        key = tt.field_key(res["title"], props)
        spec = tt.spec_for(kind, description=res["description"], options=res["options"])
        if res["title"] != tt.title_of(key, None):
            spec = {"title": res["title"], **spec}
        props[key] = spec
        if res["required"]:
            self.schema.setdefault("required", []).append(key)
        return key

    def delete_field(self, key: str) -> None:
        n = self.placements().get(key, 0)
        title = tt.title_of(key, self.schema["properties"].get(key))
        msg = t("Eliminare il campo «{n}»?", n=title)
        if n:
            msg += "\n\n" + t("Verrà tolto anche dal documento.")
        if not messagebox.askyesno(t("Elimina campo"), msg, parent=self):
            return
        self.snapshot()
        self.unplace_all(key, record=False)
        self.schema["properties"].pop(key, None)
        self.schema["required"] = [r for r in self.schema.get("required", []) if r != key]
        if self.selected == key:
            self.selected = None
        self.changed()

    def unplace_all(self, key: str, *, record: bool = True) -> None:
        if record:
            self.snapshot()
        self.remove_everywhere(key)
        if record:
            self.changed()

    def remove_everywhere(self, key: str) -> None:
        raise NotImplementedError

    def unplace(self, payload: Dict[str, Any]) -> None:
        raise NotImplementedError

    def name_of(self, key: str) -> str:
        return tt.title_of(key, self.schema["properties"].get(key))

    # ------------------------------------------------------------ annulla
    def doc_state(self) -> Any:
        return None

    def restore_doc_state(self, state: Any) -> None:
        pass

    def snapshot(self) -> None:
        self.undo_stack.append((copy.deepcopy(self.schema), self.doc_state()))
        del self.undo_stack[:-30]
        self.undo_btn.configure(state="normal")

    def drop_snapshot(self) -> None:
        if self.undo_stack:
            self.undo_stack.pop()
        if not self.undo_stack:
            self.undo_btn.configure(state="disabled")

    def undo(self) -> None:
        if not self.undo_stack:
            return
        schema, state = self.undo_stack.pop()
        self.schema = schema
        self.restore_doc_state(state)
        if self.selected not in self.schema.get("properties", {}):
            self.selected = None
        self.dirty = True
        if not self.undo_stack:
            self.undo_btn.configure(state="disabled")
        self.refresh()

    def changed(self) -> None:
        self.dirty = True
        self.refresh()

    def refresh(self) -> None:
        self.render_canvas()
        self.render_fields()
        self.render_props()
        self.highlight_selected()

    def render_canvas(self) -> None:
        pass

    def _clean_schema(self) -> Dict[str, Any]:
        sch = copy.deepcopy(self.schema)
        props = sch.get("properties", {})
        sch["required"] = [r for r in dict.fromkeys(sch.get("required", [])) if r in props]
        sch.setdefault("type", "object")
        return sch

    def _welcome(self) -> None:
        if self.first_time:
            self.banner.configure(text="  " + t("Ora trascina i campi nei punti del documento che l'AI deve "
                                                "compilare.") + "  ")
            self.banner.pack(side="left", padx=12)

    def _hide_welcome(self) -> None:
        if self.first_time and self.armed is None:
            self.first_time = False
            self.banner.pack_forget()


def _short(s: str, n: int) -> str:
    return s if len(s) <= n else s[: n - 1] + "…"


# ====================================================================== Word
class TemplateEditor(_VisualEditor):
    CANVAS_HINT = ""

    def __init__(self, win: Any, mod, *, first_time: bool = False):
        import docx
        self.path = Path(mod.template_path)
        self.doc = docx.Document(str(self.path))
        self.segments: List[Dict[str, Any]] = []
        self.CANVAS_HINT = t("Rilasciala sulla riga da compilare (es. Nome: ______), in una cella vuota "
                             "o su un testo selezionato.")
        super().__init__(win, mod, t("Editor visuale · {name}", name=mod.name), first_time=first_time)
        self.view_mode = ctk.CTkSegmentedButton(self.toolbar, values=[t("Modifica"), t("Anteprima con dati")],
                                                command=lambda v: self._mode_changed())
        self.view_mode.set(t("Modifica"))
        self.view_mode.pack(side="right")
        self._build_view()
        self.refresh()
        self._welcome()

    def _build_view(self) -> None:
        """Vista semplificata (testo e tabelle): usata solo senza Word/LibreOffice."""
        box = self.canvas_box
        box.rowconfigure(0, weight=1)
        box.columnconfigure(0, weight=1)
        self.text = tk.Text(box, wrap="word", relief="flat", bd=0, padx=36, pady=28, cursor="arrow",
                            font=("Calibri", 12), bg=col("surface"), fg=col("text"), highlightthickness=1,
                            highlightbackground=col("border"), insertbackground=col("primary"),
                            insertwidth=3, spacing1=2, spacing3=4, selectbackground=col("selection"),
                            selectforeground=col("text"), inactiveselectbackground=col("selection"))
        self.text.grid(row=0, column=0, sticky="nsew")
        sb = ctk.CTkScrollbar(box, command=self.text.yview)
        sb.grid(row=0, column=1, sticky="ns")
        self.text.configure(yscrollcommand=sb.set)
        tx = self.text
        tx.tag_configure("ph", background=col("primary_soft"), foreground=col("link"), relief="raised",
                         borderwidth=1, font=("Calibri", 12, "bold"))
        tx.tag_configure("ph_sel", background=col("primary"), foreground="#ffffff")
        tx.tag_configure("brace", elide=True)
        tx.tag_configure("h1", font=("Calibri", 18, "bold"), spacing1=10, spacing3=6)
        tx.tag_configure("h2", font=("Calibri", 15, "bold"), spacing1=8, spacing3=4)
        tx.tag_configure("cell", lmargin1=8, lmargin2=8)
        tx.tag_configure("f_b", font=("Calibri", 12, "bold"))
        tx.tag_configure("f_i", font=("Calibri", 12, "italic"))
        tx.tag_configure("f_bi", font=("Calibri", 12, "bold italic"))
        tx.tag_configure("u", underline=True)
        tx.tag_configure("al_center", justify="center")
        tx.tag_configure("al_right", justify="right")
        tx.tag_raise("ph")
        tx.tag_configure("sep", foreground=col("border"))
        tx.tag_configure("empty", foreground=col("text_faint"), background=col("surface_alt"))
        tx.tag_configure("blank", foreground=col("text_faint"))
        tx.tag_configure("append", foreground=col("text_faint"), font=("Calibri", 11, "italic"),
                         justify="center", spacing1=18)
        tx.tag_configure("drop", background=col("success_soft"), foreground=col("success"), relief="solid",
                         borderwidth=1)
        tx.tag_raise("sel")
        tx.tag_raise("drop")
        tx.bind("<Key>", self._key)
        for ev in ("<<Paste>>", "<<PasteSelection>>", "<<Cut>>", "<<Clear>>", "<Button-2>"):
            tx.bind(ev, lambda e: "break")
        tx.bind("<ButtonPress-1>", self._press)
        tx.bind("<B1-Motion>", self.drag.motion)
        tx.bind("<ButtonRelease-1>", self.drag.release)
        tx.bind("<Motion>", self._hover_cursor)
        self.drag.add_target(tx, hover=self._hover, drop=self._drop, leave=self._clear_drop)

    def _footer_extra(self) -> None:
        button(self.footer, t("Apri copia in Word"), self.open_word, icon_name="external-link").pack(
            side="right", pady=12)

    def tips(self) -> List[str]:
        return [t("Seleziona un testo d'esempio (es. «Mario Rossi») e trascinaci sopra un campo: lo sostituisce."),
                t("Trascina un campo in fondo al documento per aggiungere una nuova riga «Nome: campo»."),
                *super().tips()]

    # ------------------------------------------------------------ dati
    def placements(self) -> Dict[str, int]:
        out: Dict[str, int] = {}
        for _p, _a, _b, k in tt.iter_placeholders(_iter_all_paragraphs(self.doc)):
            out[k] = out.get(k, 0) + 1
        return out

    def doc_state(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()

    def restore_doc_state(self, state: bytes) -> None:
        import docx
        self.doc = docx.Document(io.BytesIO(state))

    def remove_everywhere(self, key: str) -> None:
        tt.remove_field(_iter_all_paragraphs(self.doc), key)

    def unplace(self, payload: Dict[str, Any]) -> None:
        p = payload["p"]
        tt.remove_field_at(p, payload["start"], payload["end"])

    def where_text(self, key: str) -> str:
        n = self.placements().get(key, 0)
        if not n:
            return t("Non ancora nel documento: trascinalo dall'elenco a sinistra nel punto giusto.")
        ph = "{{" + key + "}}"
        return t("Nel documento come {ph} · {n} volte", ph=ph, n=n) if n > 1 else \
            t("Nel documento come {ph}", ph=ph)

    # ------------------------------------------------------------ disegno del documento
    def _preview(self) -> bool:
        return self.view_mode.get() == t("Anteprima con dati") if hasattr(self, "view_mode") else False

    def _mode_changed(self) -> None:
        self.disarm()
        self.render_canvas()

    def _sample_values(self) -> Dict[str, Any]:
        rows, _n = self.win.db.search_reports(module_id=self.mod.id, limit=1)
        if rows:
            import json
            try:
                return json.loads(rows[0].get("final_json") or rows[0].get("draft_json") or "{}")
            except ValueError:
                pass
        return {k: f"«{tt.title_of(k, s)}»" for k, s in self.schema.get("properties", {}).items()}

    def render_canvas(self) -> None:
        if not hasattr(self, "text"):
            return
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = self.doc
        preview = self._preview()
        if preview:
            from ...parsers.docx_parser import apply_placeholders
            import docx
            doc = docx.Document(io.BytesIO(self.doc_state()))
            apply_placeholders(doc, self._sample_values())
        tx = self.text
        y = tx.yview()[0]
        tx.configure(state="normal")
        tx.delete("1.0", "end")
        self.segments = []

        def add_par(p, tags=(), hint: str = "") -> None:
            start = tx.index("end-1c")
            if p.text:
                runs = list(p.runs)
                if "".join(r.text for r in runs) == p.text:  # grassetto/corsivo/sottolineato dei run
                    for r in runs:
                        style = ("b" if r.bold else "") + ("i" if r.italic else "")
                        extra = tuple(x for x in (style and "f_" + style, "u" if r.underline else "") if x)
                        tx.insert("end", r.text, (*tags, *extra))
                else:
                    tx.insert("end", p.text, tags)
                end = tx.index("end-1c")
                align = getattr(p.alignment, "name", "") if p.alignment is not None else ""
                if align in ("CENTER", "RIGHT") and not tags:
                    tx.tag_add("al_" + align.lower(), start, end)
                tx.tag_add("p", start, end)
            else:
                tx.insert("end", "  " + t("(vuota)") + "  " if tags else " ", (*tags, "empty") if tags else tags)
                end = tx.index("end-1c")
            self.segments.append({"p": p, "s": start, "e": end, "len": len(p.text), "hint": hint,
                                  "cell": bool(tags)})

        for block in _iter_block_items(doc):
            if isinstance(block, Paragraph):
                style = (block.style.name if block.style is not None else "") or ""
                tag = "h1" if style in ("Title", "Titolo", "Heading 1", "Titolo 1") else (
                    "h2" if style.startswith(("Heading", "Titolo")) else None)
                line_start = tx.index("end-1c")
                add_par(block)
                if tag:
                    tx.tag_add(tag, line_start, "end-1c")
                tx.insert("end", "\n")
            elif isinstance(block, Table):
                seen = set()  # celle unite in verticale: python-docx le ripete in ogni riga
                for r in block.rows:
                    left = ""
                    first = True
                    for cell in r.cells:
                        if id(cell._tc) in seen:
                            continue
                        seen.add(id(cell._tc))
                        if not first:
                            tx.insert("end", "   │   ", "sep")
                        first = False
                        for i, p in enumerate(cell.paragraphs):
                            if i:
                                tx.insert("end", "  /  ", "sep")
                            add_par(p, ("cell",), hint=left)
                        left = cell.text.strip() or left
                    tx.insert("end", "\n")
                tx.insert("end", "\n")
        if not preview:
            tx.insert("end", "\n＋  " + t("Rilascia qui per aggiungere una nuova riga in fondo") + "\n", "append")
        text_all = tx.get("1.0", "end")
        for m in tt.PLACEHOLDER_RE.finditer(text_all):
            tx.tag_add("ph", f"1.0+{m.start()}c", f"1.0+{m.end()}c")
            tx.tag_add("brace", f"1.0+{m.start()}c", f"1.0+{m.start(1)}c")
            tx.tag_add("brace", f"1.0+{m.end(1)}c", f"1.0+{m.end()}c")
        for m in tt.BLANK_RE.finditer(text_all):
            tx.tag_add("blank", f"1.0+{m.start()}c", f"1.0+{m.end()}c")
        tx.yview_moveto(y)
        self.highlight_selected()

    def highlight_selected(self) -> None:
        if not hasattr(self, "text"):
            return
        tx = self.text
        tx.tag_remove("ph_sel", "1.0", "end")
        if self.selected:
            for m in tt.PLACEHOLDER_RE.finditer(tx.get("1.0", "end")):
                if m.group(1) == self.selected:
                    tx.tag_add("ph_sel", f"1.0+{m.start()}c", f"1.0+{m.end()}c")
        tx.tag_raise("ph_sel")

    # ------------------------------------------------------------ posizione sotto il puntatore
    def _index_at(self, x_root: int, y_root: int) -> Tuple[str, int, int]:
        tx = self.text
        x, y = x_root - tx.winfo_rootx(), y_root - tx.winfo_rooty()
        idx = tx.index(f"@{x},{y}")
        bbox = tx.bbox(idx)
        if bbox and x > bbox[0] + bbox[2] / 2 and tx.compare(idx, "<", f"{idx} lineend"):
            idx = tx.index(f"{idx}+1c")
        return idx, x, y

    def _segment_at(self, idx: str) -> Optional[int]:
        tx = self.text
        best = None
        line = idx.split(".")[0]
        for i, s in enumerate(self.segments):
            if tx.compare(s["s"], "<=", idx) and tx.compare(idx, "<=", s["e"]):
                return i
            if s["s"].split(".")[0] == line and tx.compare(s["s"], "<=", idx):
                best = i  # sul separatore di una tabella: la cella a sinistra
        return best

    def locate(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        """Dove finirebbe il campo rilasciato in (x_root, y_root)."""
        tx = self.text
        idx, _x, y = self._index_at(x_root, y_root)
        last = tx.dlineinfo("end-1c")
        if "append" in tx.tag_names(idx) or (last and y > last[1] + last[3]):
            return {"mode": "append", "hl": ("append.first", "append.last"), "suggest": ""}
        i = self._segment_at(idx)
        if i is None:
            return None
        seg = self.segments[i]
        ptext = seg["p"].text
        if not ptext:
            return {"mode": "insert", "seg": i, "start": 0, "end": None, "hl": (seg["s"], seg["e"]),
                    "suggest": tt.suggest_label(seg["hint"])}
        off = len(tx.get(seg["s"], idx)) if tx.compare(idx, ">", seg["s"]) else 0
        off = min(off, len(ptext))
        try:
            s1, s2 = tx.index("sel.first"), tx.index("sel.last")
            if (tx.compare(s1, "<=", idx) and tx.compare(idx, "<=", s2) and tx.compare(s1, ">=", seg["s"])
                    and tx.compare(s2, "<=", seg["e"])):
                a, b = len(tx.get(seg["s"], s1)), len(tx.get(seg["s"], s2))
                if b > a:
                    return {"mode": "replace", "seg": i, "start": a, "end": b, "hl": (s1, s2),
                            "suggest": tt.suggest_label(ptext[:a]) or tt.suggest_label(seg["hint"])}
        except tk.TclError:
            pass
        ph = tt.placeholder_at(ptext, off) or (tt.placeholder_at(ptext, off - 1) if off else None)
        if ph:
            return {"mode": "replace", "seg": i, "start": ph[0], "end": ph[1], "on_key": ph[2],
                    "hl": (f"{seg['s']}+{ph[0]}c", f"{seg['s']}+{ph[1]}c"), "suggest": ""}
        blank = tt.blank_at(ptext, off)
        if blank:
            return {"mode": "replace", "seg": i, "start": blank[0], "end": blank[1],
                    "hl": (f"{seg['s']}+{blank[0]}c", f"{seg['s']}+{blank[1]}c"),
                    "suggest": tt.suggest_label(ptext[:blank[0]]) or tt.suggest_label(seg["hint"])}
        off = tt.snap_offset(ptext, off)
        return {"mode": "insert", "seg": i, "start": off, "end": None, "caret": f"{seg['s']}+{off}c",
                "suggest": tt.suggest_label(ptext[:off]) or tt.suggest_label(seg["hint"])}

    # ------------------------------------------------------------ eventi
    def _key(self, e) -> Optional[str]:
        if e.keysym in ("Left", "Right", "Up", "Down", "Prior", "Next", "Home", "End"):
            return None
        if e.keysym.lower() in ("c", "a") and e.state & 0x4:
            return None
        if e.keysym.lower() == "z" and e.state & 0x4:
            self.undo()
        return "break"

    def _occurrence_at(self, idx: str) -> Optional[Dict[str, Any]]:
        if "ph" not in self.text.tag_names(idx):
            return None
        i = self._segment_at(idx)
        if i is None:
            return None
        seg = self.segments[i]
        off = len(self.text.get(seg["s"], idx))
        ph = tt.placeholder_at(seg["p"].text, off)
        if not ph:
            return None
        return {"kind": "placed", "key": ph[2], "p": seg["p"], "start": ph[0], "end": ph[1], "seg": i}

    def _press(self, e) -> Optional[str]:
        if self._preview():
            return None
        if self.armed is not None:
            payload = self.armed
            self.disarm()
            self.after_idle(lambda: self._drop(e.x_root, e.y_root, payload))
            return "break"
        idx = self.text.index(f"@{e.x},{e.y}")
        occ = self._occurrence_at(idx)
        if occ is None:
            self.text.focus_set()
            return None
        if occ["key"] in self.schema["properties"]:
            click = (lambda k=occ["key"]: self.after_idle(lambda: self.select(k)))
        else:
            click = None
        self.drag.press(e, occ, self.name_of(occ["key"]), click)
        return "break"

    def _hover_cursor(self, e) -> None:
        if self.drag.dragging:
            return
        if self.armed is not None:
            cur = "crosshair"
        elif not self._preview() and "ph" in self.text.tag_names(self.text.index(f"@{e.x},{e.y}")):
            cur = "hand2"
        else:
            cur = "xterm"
        if str(self.text.cget("cursor")) != cur:
            self.text.configure(cursor=cur)

    def on_arm(self, on: bool) -> None:
        if hasattr(self, "text"):
            self.text.configure(cursor="crosshair" if on else "xterm")

    def _clear_drop(self) -> None:
        self.text.tag_remove("drop", "1.0", "end")

    def _hover(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> bool:
        self._clear_drop()
        if self._preview():
            return False
        tx = self.text
        _idx, _x, y = self._index_at(x_root, y_root)
        h = tx.winfo_height()
        if y < 30:
            tx.yview_scroll(-1, "units")
        elif y > h - 30:
            tx.yview_scroll(1, "units")
        loc = self.locate(x_root, y_root, payload)
        if loc is None:
            return False
        if "hl" in loc:
            tx.tag_add("drop", *loc["hl"])
        if "caret" in loc:
            tx.mark_set("insert", loc["caret"])
            tx.focus_set()
        return True

    def _drop(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> None:
        if self._preview():
            self.win.toast(t("Passa a «Modifica» per spostare i campi."), "info")
            return
        loc = self.locate(x_root, y_root, payload)
        if loc is None:
            self.win.toast(t("Rilascia il campo su una riga del documento."), "warning")
            return
        kind = payload.get("kind")
        if kind == "placed" and loc.get("seg") == payload.get("seg") and loc["mode"] != "append" and \
                payload["start"] <= loc["start"] <= payload["end"]:
            return  # rilasciato su se stesso
        self.snapshot()
        if kind == "new":
            key = self.create_field(payload["type"], loc.get("suggest", ""))
            if key is None:
                self.drop_snapshot()
                return
        else:
            key = payload["key"]
        if kind == "placed":
            p0 = payload["p"]
            before = len(p0.text)
            tt.remove_field_at(p0, payload["start"], payload["end"])
            delta = before - len(p0.text)
            if loc["mode"] != "append" and self.segments[loc["seg"]]["p"] is p0 and loc["start"] >= payload["end"]:
                loc["start"] -= delta
                if loc.get("end") is not None:
                    loc["end"] -= delta
        if loc["mode"] == "append":
            p = self.doc.add_paragraph()
            p.add_run(self.name_of(key) + ": ")
            tt.place_field(p, len(p.text), key)
        else:
            p = self.segments[loc["seg"]]["p"]
            if loc["mode"] == "replace":
                tt.place_field(p, loc["start"], key, end=loc["end"])
            else:
                tt.place_field(p, loc["start"], key)
        self.selected = key
        self.changed()
        self._hide_welcome()

    # ------------------------------------------------------------ salvataggio
    def save(self) -> None:
        if not self.dirty:
            self.close()
            return
        backup_dir = self.mod.folder_path / "_versions"
        backup_dir.mkdir(exist_ok=True)
        shutil.copy2(self.path, backup_dir / f"template_{time.strftime('%Y%m%d_%H%M%S')}{self.path.suffix}")
        try:
            self.doc.save(str(self.path))
        except PermissionError:
            messagebox.showerror(t("Salvataggio non riuscito"),
                                 t("Il file del template è aperto in un altro programma (es. Word). Chiudilo e "
                                   "riprova."), parent=self)
            return
        schema = self._clean_schema()
        if schema != (self.mod.schema or {}):
            self.win.mm.bump_version(self.mod.slug, new_schema=schema)
        self.win.refresh_modules(quiet=True)
        self.win.toast(t("Modulo salvato (copia precedente in _versions)."), "success")
        self.close()

    def open_word(self) -> None:
        """Apre in Word una COPIA del template con le modifiche correnti (solo da consultare)."""
        tmp = Path(tempfile.mkdtemp(prefix="mai_tpl_")) / self.path.name
        self.doc.save(str(tmp))
        os.startfile(str(tmp))  # type: ignore[attr-defined]
        self.win.toast(t("Aperta una copia di consultazione: le modifiche fatte in Word non vengono salvate nel "
                         "modulo."), "info")


class PageTemplateEditor(TemplateEditor):
    """Editor Word con le pagine vere del documento (impaginate da Word o LibreOffice).

    Il modulo si vede esattamente com'e': caratteri, tabelle, immagini, intestazioni,
    margini. Dopo ogni modifica la pagina viene impaginata di nuovo; nel frattempo
    il rilascio dei campi e' sospeso (le posizioni non sarebbero aggiornate).
    """

    GAP = 18

    def __init__(self, win: Any, mod, *, first_time: bool = False):
        self.layout: Optional[Any] = None
        self.infos: List[Any] = []
        self.pages: List[Dict[str, Any]] = []  # {"img", "x", "y", "w", "h"}
        self.scale = 1.0
        self.busy = False
        self._gen = 0
        self._pdf: Optional[Path] = None
        self._tmp = Path(tempfile.mkdtemp(prefix="docxai_edit_"))
        self._sel: Optional[Tuple[int, int, int]] = None      # (paragrafo, inizio, fine)
        self._anchor: Optional[Tuple[int, int]] = None
        self._last_width = 0
        self._resize_job: Optional[str] = None
        self.error = ""
        super().__init__(win, mod, first_time=first_time)

    # ------------------------------------------------------------ vista
    def _build_view(self) -> None:
        box = self.canvas_box
        box.rowconfigure(0, weight=1)
        box.columnconfigure(0, weight=1)
        self.cv = tk.Canvas(box, bg=col("surface_alt"), highlightthickness=1, highlightbackground=col("border"),
                            cursor="arrow")
        self.cv.grid(row=0, column=0, sticky="nsew")
        vs = ctk.CTkScrollbar(box, command=self.cv.yview)
        vs.grid(row=0, column=1, sticky="ns")
        hs = ctk.CTkScrollbar(box, command=self.cv.xview, orientation="horizontal")
        hs.grid(row=1, column=0, sticky="ew")
        self.cv.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.cv.bind("<ButtonPress-1>", self._press)
        self.cv.bind("<B1-Motion>", self._motion)
        self.cv.bind("<ButtonRelease-1>", self._release)
        self.cv.bind("<Motion>", self._hover_cursor)
        self.cv.bind("<MouseWheel>", lambda e: self.cv.yview_scroll(-1 if e.delta > 0 else 1, "units"))
        self.cv.bind("<Button-4>", lambda e: self.cv.yview_scroll(-1, "units"))
        self.cv.bind("<Button-5>", lambda e: self.cv.yview_scroll(1, "units"))
        self.cv.bind("<Configure>", self._on_resize)
        self.cv.bind("<Destroy>", lambda e: shutil.rmtree(self._tmp, ignore_errors=True)
                     if e.widget is self.cv else None)
        self.drag.add_target(self.cv, hover=self._hover, drop=self._drop, leave=self._clear_drop)
        self.text = None  # type: ignore[assignment] - la vista testuale non esiste

    def tips(self) -> List[str]:
        return [t("Il modulo è mostrato esattamente come in Word: i campi sono evidenziati in azzurro."),
                t("Seleziona col mouse un testo d'esempio (es. «Mario Rossi») e trascinaci sopra un campo: "
                  "lo sostituisce."),
                *_VisualEditor.tips(self)]

    # ------------------------------------------------------------ impaginazione (in background)
    def render_canvas(self) -> None:
        if not hasattr(self, "cv"):
            return
        from ... import docx_layout as dl
        self._gen += 1
        gen = self._gen
        self.busy = True
        self._sel = None
        self._anchor = None
        self._draw_busy()
        preview = self._preview()
        try:
            if preview:
                import docx
                from ...parsers.docx_parser import apply_placeholders
                doc = docx.Document(io.BytesIO(self.doc_state()))
                apply_placeholders(doc, self._sample_values())
                buf = io.BytesIO()
                doc.save(buf)
                data, texts = buf.getvalue(), []
            else:
                self.infos = dl.paragraphs(self.doc)
                self.segments = [{"p": i.paragraph} for i in self.infos]
                data, texts = dl.render_copy(self.doc), dl.snapshot(self.infos)
        except Exception as exc:  # noqa: BLE001
            LOG.exception("preparazione anteprima")
            self._render_failed(str(exc))
            return
        self.cv.update_idletasks()
        width = max(400, self.cv.winfo_width())
        colors = (col("primary"), col("primary_soft"))
        threading.Thread(target=self._render_worker, args=(gen, data, texts, preview, width, colors),
                         daemon=True, name="docx-render").start()

    def _render_worker(self, gen: int, data: bytes, texts: List[Any], preview: bool, width: int,
                       colors: Tuple[str, str]) -> None:
        from ... import docx_layout as dl
        try:
            src = self._tmp / f"modulo_{gen}.docx"
            pdf = self._tmp / f"modulo_{gen}.pdf"
            src.write_bytes(data)
            dl.converter().convert(src, pdf)
            layout = None if preview else dl.build_layout(pdf, texts)
            pages, scale = self._rasterize(pdf, width, layout, texts, colors)
            src.unlink(missing_ok=True)
        except Exception as exc:  # noqa: BLE001
            LOG.warning("Impaginazione del modulo non riuscita: %s", exc)
            err = str(exc)
            self._post(lambda: self._render_failed(err) if gen == self._gen else None)
            return
        self._post(lambda: self._render_done(gen, pdf, layout, pages, scale))

    def _post(self, fn) -> None:
        def run() -> None:
            # l'impaginazione puo' finire dopo la chiusura dell'editor (o il passaggio
            # alla vista semplificata): il canvas non esiste piu'
            try:
                if not (self.winfo_exists() and self.cv.winfo_exists()):
                    return
            except tk.TclError:
                return
            fn()
        try:
            self.after(0, run)
        except (RuntimeError, tk.TclError):
            pass  # finestra chiusa nel frattempo

    @staticmethod
    def _rasterize(pdf: Path, width: int, layout: Any, texts: List[Any],
                   colors: Tuple[str, str]) -> Tuple[List[Any], float]:
        import pypdfium2 as pdfium
        from PIL import Image, ImageDraw
        from ... import docx_layout as dl
        doc = pdfium.PdfDocument(str(pdf))
        try:
            pw = max((doc[i].get_size()[0] for i in range(len(doc))), default=595.0)
            scale = max(0.6, min(2.5, (width - 2 * PageTemplateEditor.GAP - 8) / pw))
            pages = []
            for i in range(len(doc)):
                img = doc[i].render(scale=scale).to_pil().convert("RGB")
                pages.append(img)
        finally:
            doc.close()
        if layout is not None:  # campi {{...}} evidenziati sulla pagina (riempimento semitrasparente)
            fill = Image.new("RGBA", (1, 1), colors[0]).getpixel((0, 0))[:3] + (60,)
            line = Image.new("RGBA", (1, 1), colors[0]).getpixel((0, 0))[:3] + (200,)
            overlays: Dict[int, Any] = {}
            for idx, (ptext, _story, _cell) in enumerate(texts):
                pl = layout.paras.get(idx)
                if pl is None:
                    continue
                for m in tt.PLACEHOLDER_RE.finditer(ptext):
                    for (x0, y0, x1, y1) in dl.span_boxes(pl, m.start(), m.end()):
                        ov = overlays.get(pl.page)
                        if ov is None:
                            ov = overlays[pl.page] = Image.new("RGBA", pages[pl.page].size, (0, 0, 0, 0))
                        ImageDraw.Draw(ov).rounded_rectangle(
                            (x0 * scale - 2, y0 * scale - 2, x1 * scale + 2, y1 * scale + 2), radius=3,
                            fill=fill, outline=line, width=1)
            for pi, ov in overlays.items():
                pages[pi] = Image.alpha_composite(pages[pi].convert("RGBA"), ov).convert("RGB")
        return pages, scale

    def _render_done(self, gen: int, pdf: Path, layout: Any, pages: List[Any], scale: float) -> None:
        if gen != self._gen or not self.cv.winfo_exists():
            return
        from PIL import ImageTk
        old, self._pdf = self._pdf, pdf
        if old is not None and old != pdf:
            old.unlink(missing_ok=True)
        self.layout = layout
        self.scale = scale
        self.busy = False
        self.error = ""
        cv = self.cv
        y_view = cv.yview()[0]
        cv.delete("all")
        self.pages = []
        width = max(cv.winfo_width(), 400)
        y = self.GAP
        max_w = 0
        for img in pages:
            photo = ImageTk.PhotoImage(img)
            w, h = img.size
            x = max(self.GAP, (width - w) // 2)
            cv.create_rectangle(x + 3, y + 3, x + w + 3, y + h + 3, fill=col("border"), outline="")
            cv.create_image(x, y, image=photo, anchor="nw")
            self.pages.append({"img": photo, "x": x, "y": y, "w": w, "h": h})
            y += h + self.GAP
            max_w = max(max_w, x + w + self.GAP)
        if not self._preview():
            cv.create_text(width // 2, y + 4, anchor="n", tags=("append",), fill=col("text_faint"),
                           font=font("small"),
                           text="＋  " + t("Rilascia qui per aggiungere una nuova riga in fondo"))
            y += 40
        cv.configure(scrollregion=(0, 0, max(max_w, width), y))
        cv.yview_moveto(y_view)
        self._last_width = width
        self.highlight_selected()
        # la finestra ha cambiato dimensione mentre si impaginava: pagina adattata alla larghezza
        if abs(max(cv.winfo_width(), 400) - (max(p["w"] for p in self.pages) + 2 * self.GAP + 8)) > 60 \
                and self._resize_job is None:
            self._resize_job = self.after(200, self._rerasterize)

    def _draw_busy(self) -> None:
        cv = self.cv
        cv.delete("busy")
        if not self.pages:
            cv.create_text(max(cv.winfo_width(), 400) // 2, 80, tags=("busy",), fill=col("text_muted"),
                           font=font("body"), text=t("Impaginazione del modulo in corso…"))
            return
        x = cv.canvasx(cv.winfo_width() // 2)
        y = cv.canvasy(14)
        cv.create_rectangle(x - 150, y, x + 150, y + 30, fill=col("info_soft"), outline=col("border"),
                            tags=("busy",))
        cv.create_text(x, y + 15, tags=("busy",), fill=col("link"), font=font("small_b"),
                       text=t("Aggiornamento della pagina…"))

    def _render_failed(self, err: str) -> None:
        from ... import docx_layout as dl
        self.busy = False
        self.error = err
        if not dl.converter().available:  # nessun motore utilizzabile: vista semplificata
            self.win.toast(t("Impaginazione non disponibile su questo PC: uso la vista semplificata."),
                           "warning")
            self.after(0, self._to_simple)
            return
        cv = self.cv
        cv.delete("all")
        self.pages = []
        w = max(cv.winfo_width(), 400)
        cv.create_text(w // 2, 70, width=w - 80, fill=col("text"), font=font("body"), justify="center",
                       text=t("Non è stato possibile impaginare il modulo con {e}.", e=_engine_name())
                       + "\n" + _short(err, 300))
        frame = ctk.CTkFrame(cv, fg_color="transparent")
        button(frame, t("Riprova"), self.render_canvas, icon_name="refresh-cw").pack(side="left", padx=6)
        button(frame, t("Usa la vista semplificata"), self._to_simple, kind="secondary").pack(side="left", padx=6)
        cv.create_window(w // 2, 140, window=frame, anchor="n")

    def _to_simple(self) -> None:
        """Riapre l'editor nella vista semplificata, conservando le modifiche in corso."""
        state, schema, dirty, selected = self.doc_state(), copy.deepcopy(self.schema), self.dirty, self.selected
        win, mod = self.win, self.mod
        self.dirty = False
        self.close()
        ed = TemplateEditor(win, mod)
        ed.restore_doc_state(state)
        ed.schema, ed.dirty, ed.selected = schema, dirty, selected
        ed.refresh()

    def _on_resize(self, e) -> None:
        if abs(e.width - self._last_width) < 40 or not self.pages:
            return
        if self._resize_job is not None:
            self.after_cancel(self._resize_job)
        self._resize_job = self.after(350, self._rerasterize)

    def _rerasterize(self) -> None:
        self._resize_job = None
        if self.busy or self._pdf is None or not self._pdf.is_file():
            return
        from ... import docx_layout as dl
        self._gen += 1
        gen, pdf, layout = self._gen, self._pdf, self.layout
        texts = dl.snapshot(self.infos) if layout is not None else []
        width = max(400, self.cv.winfo_width())
        colors = (col("primary"), col("primary_soft"))
        self.busy = True

        def work() -> None:
            try:
                pages, scale = self._rasterize(pdf, width, layout, texts, colors)
            except Exception as exc:  # noqa: BLE001
                err = str(exc)
                self._post(lambda: self._render_failed(err))
                return
            self._post(lambda: self._render_done(gen, pdf, layout, pages, scale))
        threading.Thread(target=work, daemon=True, name="docx-raster").start()

    # ------------------------------------------------------------ coordinate
    def _to_page(self, x_root: int, y_root: int) -> Optional[Tuple[int, float, float]]:
        """(pagina, x, y in punti) sotto il puntatore; pagina -1 = sotto l'ultima pagina."""
        cx = self.cv.canvasx(x_root - self.cv.winfo_rootx())
        cy = self.cv.canvasy(y_root - self.cv.winfo_rooty())
        for i, pg in enumerate(self.pages):
            if pg["y"] <= cy < pg["y"] + pg["h"] and pg["x"] <= cx <= pg["x"] + pg["w"]:
                return i, (cx - pg["x"]) / self.scale, (cy - pg["y"]) / self.scale
        if self.pages and cy >= self.pages[-1]["y"] + self.pages[-1]["h"]:
            return -1, 0.0, 0.0
        return None

    def _rect(self, page: int, box: Tuple[float, float, float, float], pad: float = 2) -> Tuple[float, ...]:
        pg = self.pages[page]
        s = self.scale
        return (pg["x"] + box[0] * s - pad, pg["y"] + box[1] * s - pad,
                pg["x"] + box[2] * s + pad, pg["y"] + box[3] * s + pad)

    def _hit(self, x_root: int, y_root: int) -> Optional[Tuple[int, int]]:
        from ... import docx_layout as dl
        if self.layout is None or self.busy:
            return None
        pos = self._to_page(x_root, y_root)
        if pos is None or pos[0] < 0:
            return None
        return dl.locate(self.layout, *pos)

    # ------------------------------------------------------------ dove finirebbe il campo
    def locate(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        from ... import docx_layout as dl
        if self.layout is None or self.busy:
            return None
        pos = self._to_page(x_root, y_root)
        if pos is None:
            return None
        last_page, bottom = self.layout.content_bottom
        if pos[0] < 0 or (pos[0] == len(self.pages) - 1 and pos[0] >= last_page and pos[2] > bottom + 24):
            return {"mode": "append", "suggest": ""}
        hit = dl.locate(self.layout, *pos)
        if hit is None:
            return None
        i, off = hit
        info = self.infos[i]
        ptext = info.paragraph.text
        pl = self.layout.paras[i]
        if not ptext:
            return {"mode": "insert", "seg": i, "start": 0, "end": None, "area": (pl.page, pl.area),
                    "suggest": tt.suggest_label(info.hint)}
        off = min(off, len(ptext))
        if self._sel and self._sel[0] == i and self._sel[1] <= off <= self._sel[2]:
            a, b = self._sel[1], self._sel[2]
            return {"mode": "replace", "seg": i, "start": a, "end": b, "span": (i, a, b),
                    "suggest": tt.suggest_label(ptext[:a]) or tt.suggest_label(info.hint)}
        ph = tt.placeholder_at(ptext, off) or (tt.placeholder_at(ptext, off - 1) if off else None)
        if ph:
            return {"mode": "replace", "seg": i, "start": ph[0], "end": ph[1], "on_key": ph[2],
                    "span": (i, ph[0], ph[1]), "suggest": ""}
        blank = tt.blank_at(ptext, off)
        if blank:
            return {"mode": "replace", "seg": i, "start": blank[0], "end": blank[1], "span": (i, *blank[:2]),
                    "suggest": tt.suggest_label(ptext[:blank[0]]) or tt.suggest_label(info.hint)}
        off = tt.snap_offset(ptext, off)
        return {"mode": "insert", "seg": i, "start": off, "end": None, "caret": (i, off),
                "suggest": tt.suggest_label(ptext[:off]) or tt.suggest_label(info.hint)}

    def _caret_box(self, i: int, off: int) -> Optional[Tuple[int, Tuple[float, float, float, float]]]:
        pl = self.layout.paras.get(i) if self.layout else None
        if pl is None or not pl.boxes:
            return None
        if off in pl.boxes:
            b = pl.boxes[off]
            return pl.page, (b[0] - 1, b[1], b[0] + 1, b[3])
        prev = max((o for o in pl.boxes if o < off), default=None)
        b = pl.boxes[prev] if prev is not None else next(iter(pl.boxes.values()))
        x = b[2] + (2 if prev is not None and off > prev + 1 else 0)
        return pl.page, (x - 1, b[1], x + 1, b[3])

    # ------------------------------------------------------------ disegno di evidenziazioni
    def _clear_drop(self) -> None:
        self.cv.delete("drop")

    def _draw_span(self, span: Tuple[int, int, int], tag: str, outline: str, fill: str = "",
                   width: int = 2, stipple: str = "") -> None:
        from ... import docx_layout as dl
        i, a, b = span
        pl = self.layout.paras.get(i) if self.layout else None
        if pl is None:
            return
        for box in dl.span_boxes(pl, a, b):
            self.cv.create_rectangle(*self._rect(pl.page, box), outline=outline, width=width, fill=fill,
                                     stipple=stipple, tags=(tag,))

    def _hover(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> bool:
        self._clear_drop()
        if self._preview() or self.busy:
            return False
        cv = self.cv
        y = y_root - cv.winfo_rooty()
        if y < 30:
            cv.yview_scroll(-1, "units")
        elif y > cv.winfo_height() - 30:
            cv.yview_scroll(1, "units")
        loc = self.locate(x_root, y_root, payload)
        if loc is None:
            return False
        ok = col("success")
        if loc["mode"] == "append":
            for item in cv.find_withtag("append"):
                x1, y1, x2, y2 = cv.bbox(item)
                cv.create_rectangle(x1 - 10, y1 - 6, x2 + 10, y2 + 6, outline=ok, width=2, tags=("drop",))
        elif "span" in loc:
            self._draw_span(loc["span"], "drop", ok, width=3)
        elif "area" in loc:
            page, area = loc["area"]
            cv.create_rectangle(*self._rect(page, area, 0), outline=ok, width=3, tags=("drop",))
        elif "caret" in loc:
            cb = self._caret_box(*loc["caret"])
            if cb:
                x1, y1, x2, y2 = self._rect(cb[0], cb[1], 0)
                cv.create_rectangle(x1 - 1, y1 - 3, x2 + 1, y2 + 3, fill=ok, outline=ok, tags=("drop",))
        return True

    def highlight_selected(self) -> None:
        if not hasattr(self, "cv"):
            return
        self.cv.delete("selph")
        if not self.selected or self.layout is None or self.busy:
            return
        for i, info in enumerate(self.infos):
            for m in tt.PLACEHOLDER_RE.finditer(info.paragraph.text):
                if m.group(1) == self.selected:
                    self._draw_span((i, m.start(), m.end()), "selph", col("primary"), width=3)

    def _draw_selection(self) -> None:
        self.cv.delete("textsel")
        if self._sel and self._sel[2] > self._sel[1]:
            self._draw_span(self._sel, "textsel", col("primary"), fill=col("primary"), width=1, stipple="gray25")

    # ------------------------------------------------------------ eventi
    def _occurrence(self, x_root: int, y_root: int) -> Optional[Dict[str, Any]]:
        hit = self._hit(x_root, y_root)
        if hit is None:
            return None
        i, off = hit
        p = self.infos[i].paragraph
        ph = tt.placeholder_at(p.text, off) or (tt.placeholder_at(p.text, off - 1) if off else None)
        if not ph:
            return None
        return {"kind": "placed", "key": ph[2], "p": p, "start": ph[0], "end": ph[1], "seg": i}

    def _press(self, e) -> Optional[str]:
        self.cv.focus_set()
        if self._preview():
            return None
        if self.armed is not None:
            payload = self.armed
            self.disarm()
            self.after_idle(lambda: self._drop(e.x_root, e.y_root, payload))
            return "break"
        occ = self._occurrence(e.x_root, e.y_root)
        if occ is not None:
            click = (lambda k=occ["key"]: self.after_idle(lambda: self.select(k))) \
                if occ["key"] in self.schema["properties"] else None
            self.drag.press(e, occ, self.name_of(occ["key"]), click)
            return "break"
        hit = self._hit(e.x_root, e.y_root)
        self._anchor = hit
        self._sel = None
        self._draw_selection()
        return "break"

    def _motion(self, e) -> Optional[str]:
        if self.drag._pending is not None:
            return self.drag.motion(e)
        if self._anchor is None:
            return None
        hit = self._hit(e.x_root, e.y_root)
        if hit is None or hit[0] != self._anchor[0]:
            return "break"  # la selezione resta dentro un paragrafo
        a, b = sorted((self._anchor[1], hit[1]))
        self._sel = (hit[0], a, b) if b > a else None
        self._draw_selection()
        return "break"

    def _release(self, e) -> Optional[str]:
        self._anchor = None
        if self.drag._pending is not None:
            return self.drag.release(e)
        return None

    def _hover_cursor(self, e) -> None:
        if self.drag.dragging:
            return
        if self.armed is not None:
            cur = "crosshair"
        elif self._preview() or self.busy:
            cur = "arrow"
        elif self._occurrence(e.x_root, e.y_root) is not None:
            cur = "hand2"
        elif self._hit(e.x_root, e.y_root) is not None:
            cur = "xterm"
        else:
            cur = "arrow"
        if str(self.cv.cget("cursor")) != cur:
            self.cv.configure(cursor=cur)

    def on_arm(self, on: bool) -> None:
        if hasattr(self, "cv"):
            self.cv.configure(cursor="crosshair" if on else "arrow")

    def _drop(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> None:
        if self.busy and not self._preview():
            self.win.toast(t("Attendi l'aggiornamento della pagina, poi rilascia di nuovo il campo."), "info")
            return
        super()._drop(x_root, y_root, payload)


def _engine_name() -> str:
    from ... import docx_layout as dl
    return dl.converter().engine or "Word"


# ====================================================================== Excel
class SheetEditor(_VisualEditor):
    HEAD_W, HEAD_H = 44, 24

    def __init__(self, win: Any, mod, *, first_time: bool = False):
        from ...parsers.xlsx_parser import load_workbook_safe
        self.wb = load_workbook_safe(Path(mod.template_path))
        self.mapping: Dict[str, Any] = copy.deepcopy(mod.mapping or {})
        self.sheet = self.wb.sheetnames[0] if self.wb.sheetnames else ""
        self.model: Any = None
        self.CANVAS_HINT = t("Rilasciala sulla cella dove va scritto il valore (di solito accanto "
                             "all'etichetta).")
        super().__init__(win, mod, t("Editor visuale · {name}", name=mod.name), first_time=first_time)
        if len(self.wb.sheetnames) > 1:
            self.tabs = ctk.CTkSegmentedButton(self.toolbar, values=list(self.wb.sheetnames),
                                               command=self._sheet_changed)
            self.tabs.set(self.sheet)
            self.tabs.pack(side="right")
        box = self.canvas_box
        box.rowconfigure(0, weight=1)
        box.columnconfigure(0, weight=1)
        self.cv = tk.Canvas(box, bg=col("surface_alt"), highlightthickness=1, highlightbackground=col("border"))
        self.cv.grid(row=0, column=0, sticky="nsew")
        vs = ctk.CTkScrollbar(box, command=self.cv.yview)
        vs.grid(row=0, column=1, sticky="ns")
        hs = ctk.CTkScrollbar(box, command=self.cv.xview, orientation="horizontal")
        hs.grid(row=1, column=0, sticky="ew")
        self.cv.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.cv.bind("<ButtonPress-1>", self._press)
        self.cv.bind("<B1-Motion>", self.drag.motion)
        self.cv.bind("<ButtonRelease-1>", self.drag.release)
        self.cv.bind("<MouseWheel>", lambda e: self.cv.yview_scroll(-1 if e.delta > 0 else 1, "units"))
        self.cv.bind("<Shift-MouseWheel>", lambda e: self.cv.xview_scroll(-1 if e.delta > 0 else 1, "units"))
        self.cv.bind("<Motion>", self._hover_cursor)
        self.drag.add_target(self.cv, hover=self._hover, drop=self._drop, leave=self._clear_drop)
        self.refresh()
        self._welcome()

    def tips(self) -> List[str]:
        return [t("Le celle in verde sono libere; in arancione hanno già un testo che verrà sostituito."),
                t("Un campo «Elenco» scrive tutte le righe nella stessa cella, andando a capo."),
                *super().tips()]

    # ------------------------------------------------------------ dati
    def placements(self) -> Dict[str, int]:
        return {k: 1 for k, v in self.mapping.items() if isinstance(v, dict) and v.get("cell")}

    def doc_state(self) -> Dict[str, Any]:
        return copy.deepcopy(self.mapping)

    def restore_doc_state(self, state: Dict[str, Any]) -> None:
        self.mapping = state

    def remove_everywhere(self, key: str) -> None:
        self.mapping.pop(key, None)

    def unplace(self, payload: Dict[str, Any]) -> None:
        self.mapping.pop(payload["key"], None)

    def where_text(self, key: str) -> str:
        m = self.mapping.get(key)
        if not isinstance(m, dict) or not m.get("cell"):
            return t("Non ancora nel foglio: trascinalo dall'elenco a sinistra sulla cella giusta.")
        return t("Nel foglio «{s}», cella {c}", s=m.get("sheet") or self.sheet, c=m.get("cell"))

    def _sheet_changed(self, name: str) -> None:
        self.sheet = name
        self.render_canvas()

    # ------------------------------------------------------------ griglia
    def _ws(self):
        return self.wb[self.sheet] if self.sheet in self.wb.sheetnames else self.wb.active

    def render_canvas(self) -> None:
        """Il foglio come in Excel: larghezze/altezze reali, celle unite, colori, bordi,
        caratteri, allineamenti, a capo, immagini (sfondo bianco anche col tema scuro)."""
        if not hasattr(self, "cv"):
            return
        from openpyxl.utils import get_column_letter
        from ... import xlsx_layout as xl
        cv, ws = self.cv, self._ws()
        cv.delete("all")
        self._photos: List[Any] = []
        if not hasattr(self, "_theme"):
            self._theme = xl.theme_colors(self.wb)
        model = self.model = xl.build(ws, self._theme)
        ox, oy = self.HEAD_W, self.HEAD_H
        W, H = ox + model.width, oy + model.height
        head_bg, head_line, head_txt = "#F3F3F3", "#D4D4D4", "#555555"
        grid = "#E1E1E1"
        cv.create_rectangle(ox, oy, W, H, fill="#FFFFFF", outline="")
        # intestazioni di colonna e riga (come in Excel)
        cv.create_rectangle(0, 0, ox, oy, fill=head_bg, outline=head_line)
        for c in range(1, len(model.xs)):
            x1, x2 = ox + model.xs[c - 1], ox + model.xs[c]
            if x2 <= x1:
                continue
            cv.create_rectangle(x1, 0, x2, oy, fill=head_bg, outline=head_line)
            cv.create_text((x1 + x2) / 2, oy / 2, text=get_column_letter(c), fill=head_txt, font=("Segoe UI", 9))
        for r in range(1, len(model.ys)):
            y1, y2 = oy + model.ys[r - 1], oy + model.ys[r]
            if y2 <= y1:
                continue
            cv.create_rectangle(0, y1, ox, y2, fill=head_bg, outline=head_line)
            cv.create_text(ox / 2, (y1 + y2) / 2, text=str(r), fill=head_txt, font=("Segoe UI", 9))
        # griglia
        if model.gridlines:
            for x in sorted(set(model.xs[1:])):
                cv.create_line(ox + x, oy, ox + x, H, fill=grid)
            for y in sorted(set(model.ys[1:])):
                cv.create_line(ox, oy + y, W, oy + y, fill=grid)
        # celle unite: niente griglia all'interno
        for (r, c), _end in model.merged.items():
            x1, y1, x2, y2 = xl.cell_box(model, (r, c))
            if (r, c) not in model.cells or not model.cells[(r, c)].fill:
                cv.create_rectangle(ox + x1 + 1, oy + y1 + 1, ox + x2 - 1, oy + y2 - 1, fill="#FFFFFF", outline="")
        for d in model.cells.values():
            if d.fill:
                x1, y1, x2, y2 = d.box
                cv.create_rectangle(ox + x1, oy + y1, ox + x2, oy + y2, fill=d.fill, outline="")
        by_cell = {(v.get("sheet") or self.sheet, str(v.get("cell")).upper()): k
                   for k, v in self.mapping.items() if isinstance(v, dict) and v.get("cell")}
        placed = {}
        for (sheet, ref), key in by_cell.items():
            if sheet == self.sheet:
                rc = self._rc(ref)
                if rc:
                    placed[model.anchor.get(rc, rc)] = key
        for (r, c), d in model.cells.items():
            if d.text and (r, c) not in placed:
                self._draw_cell_text(d, ox, oy)
        for d in model.cells.values():
            x1, y1, x2, y2 = d.box
            for side, (width, color, dash) in d.borders.items():
                pts = {"left": (x1, y1, x1, y2), "right": (x2, y1, x2, y2), "top": (x1, y1, x2, y1),
                       "bottom": (x1, y2, x2, y2)}[side]
                cv.create_line(ox + pts[0], oy + pts[1], ox + pts[2], oy + pts[3], fill=color, width=width,
                               dash=dash or None)
        self._draw_images(model, ox, oy)
        for rc, key in placed.items():
            x1, y1, x2, y2 = xl.cell_box(model, rc)
            if x2 <= x1 or y2 <= y1:
                continue
            on = key == self.selected
            cv.create_rectangle(ox + x1 + 2, oy + y1 + 2, ox + x2 - 2, oy + y2 - 2, outline=col("primary"),
                                width=2 if on else 1, fill=col("primary") if on else "#E8EFFF",
                                tags=("pill", f"k:{key}"))
            cv.create_text(ox + x1 + 6, oy + (y1 + y2) / 2, anchor="w", text="▣ " + self.name_of(key),
                           fill="#ffffff" if on else "#1D4ED8", font=font("small_b"),
                           width=max(20, x2 - x1 - 10), tags=("pill", f"k:{key}"))
        cv.configure(scrollregion=(0, 0, W + 20, H + 20))

    def _font(self, spec: Tuple[str, int, bool, bool, bool, bool]):
        import tkinter.font as tkfont
        cache = self.__dict__.setdefault("_fonts", {})
        f = cache.get(spec)
        if f is None:
            name, size, bold, italic, under, strike = spec
            f = cache[spec] = tkfont.Font(family=name, size=-max(6, round(size * 96 / 72)),
                                          weight="bold" if bold else "normal",
                                          slant="italic" if italic else "roman", underline=under,
                                          overstrike=strike)
        return f

    @staticmethod
    def _fit(text: str, f, width: int) -> str:
        """Il testo che entra in ``width`` pixel (Excel taglia, senza puntini)."""
        if width <= 0:
            return ""
        if f.measure(text) <= width:
            return text
        lo, hi = 0, len(text)
        while lo < hi:
            mid = (lo + hi + 1) // 2
            if f.measure(text[:mid]) <= width:
                lo = mid
            else:
                hi = mid - 1
        return text[:lo]

    def _wrap(self, text: str, f, width: int, max_h: int) -> List[str]:
        lines: List[str] = []
        line_h = f.metrics("linespace")
        for para in text.split("\n"):
            cur = ""
            for word in para.split(" "):
                cand = word if not cur else cur + " " + word
                if f.measure(cand) <= width or not cur:
                    cur = cand if f.measure(cand) <= width else self._fit(cand, f, width)
                else:
                    lines.append(cur)
                    cur = self._fit(word, f, width)
            lines.append(cur)
        return lines[:max(1, max_h // max(1, line_h))]

    def _draw_cell_text(self, d: Any, ox: int, oy: int) -> None:
        x1, y1, x2, y2 = d.box
        if x2 - x1 <= 2 or y2 - y1 <= 2:
            return
        f = self._font(d.font)
        pad = 3
        right = max(x2, d.clip_right)
        if d.wrap or "\n" in d.text:
            lines = self._wrap(d.text, f, x2 - x1 - 2 * pad, y2 - y1) if d.wrap else \
                [self._fit(s, f, right - x1 - 2 * pad) for s in d.text.split("\n")][:max(
                    1, (y2 - y1) // max(1, f.metrics("linespace")))]
            text = "\n".join(lines)
        else:
            avail = (right if d.halign == "left" else x2) - x1 - 2 * pad
            text = self._fit(d.text, f, avail)
            if d.halign == "right" and text != d.text and d.text[:1].isdigit():
                text = "#" * max(1, avail // max(1, f.measure("#")))  # numero troppo lungo: ###
        anchor_x = {"left": ("w", x1 + pad), "center": ("", (x1 + x2) / 2), "right": ("e", x2 - pad)}[d.halign]
        lines_n = text.count("\n") + 1
        th = lines_n * f.metrics("linespace")
        if d.valign == "top":
            cy = y1 + 1 + th / 2
        elif d.valign == "center":
            cy = (y1 + y2) / 2
        else:
            cy = y2 - 2 - th / 2
        anchor = anchor_x[0] or "center"
        justify = {"left": "left", "center": "center", "right": "right"}[d.halign]
        self.cv.create_text(ox + anchor_x[1], oy + cy, text=text, anchor=anchor, font=f, fill=d.color,
                            justify=justify)

    def _draw_images(self, model: Any, ox: int, oy: int) -> None:
        if not model.images:
            return
        try:
            from PIL import Image, ImageTk
        except ImportError:
            return
        for x, y, w, h, data in model.images:
            try:
                img = Image.open(io.BytesIO(data)).convert("RGBA").resize((w, h))
                photo = ImageTk.PhotoImage(img)
            except Exception:  # noqa: BLE001
                continue
            self._photos.append(photo)
            self.cv.create_image(ox + x, oy + y, image=photo, anchor="nw")

    @staticmethod
    def _rc(ref: str) -> Optional[Tuple[int, int]]:
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        try:
            letters, row = coordinate_from_string(str(ref).upper())
            return int(row), column_index_from_string(letters)
        except Exception:  # noqa: BLE001
            return None

    def _cell_at(self, x_root: int, y_root: int) -> Optional[Tuple[int, int]]:
        from ... import xlsx_layout as xl
        model = getattr(self, "model", None)
        if model is None:
            return None
        x = self.cv.canvasx(x_root - self.cv.winfo_rootx()) - self.HEAD_W
        y = self.cv.canvasy(y_root - self.cv.winfo_rooty()) - self.HEAD_H
        return xl.cell_at(model, x, y)

    def _box(self, rc: Tuple[int, int]) -> Tuple[int, int, int, int]:
        from ... import xlsx_layout as xl
        x1, y1, x2, y2 = xl.cell_box(self.model, rc)
        return x1 + self.HEAD_W, y1 + self.HEAD_H, x2 + self.HEAD_W, y2 + self.HEAD_H

    def _ref(self, rc: Tuple[int, int]) -> str:
        from openpyxl.utils import get_column_letter
        return f"{get_column_letter(rc[1])}{rc[0]}"

    def _cell_text(self, rc: Tuple[int, int]) -> str:
        v = self._ws().cell(row=rc[0], column=rc[1]).value
        return "" if v is None else str(v).strip()

    def _key_at_cell(self, rc: Tuple[int, int]) -> Optional[str]:
        ref = self._ref(rc)
        for k, v in self.mapping.items():
            if isinstance(v, dict) and (v.get("sheet") or self.sheet) == self.sheet and \
                    str(v.get("cell")).upper() == ref:
                return k
        return None

    def _suggest(self, rc: Tuple[int, int]) -> str:
        r, c = rc
        for rr, cc in ((r, c - 1), (r - 1, c), (r, c - 2)):
            if rr >= 1 and cc >= 1:
                s = self._cell_text((rr, cc))
                if s and not tt.PLACEHOLDER_RE.search(s):
                    return tt.suggest_label(s)
        return ""

    # ------------------------------------------------------------ eventi
    def _press(self, e) -> Optional[str]:
        if self.armed is not None:
            payload = self.armed
            self.disarm()
            self.after_idle(lambda: self._drop(e.x_root, e.y_root, payload))
            return "break"
        rc = self._cell_at(e.x_root, e.y_root)
        key = self._key_at_cell(rc) if rc else None
        if key is None:
            return None
        click = (lambda k=key: self.after_idle(lambda: self.select(k))) if key in self.schema["properties"] else None
        self.drag.press(e, {"kind": "placed", "key": key}, self.name_of(key), click)
        return "break"

    def _hover_cursor(self, e) -> None:
        if self.drag.dragging:
            return
        rc = self._cell_at(e.x_root, e.y_root)
        cur = "crosshair" if self.armed is not None else ("hand2" if rc and self._key_at_cell(rc) else "arrow")
        if str(self.cv.cget("cursor")) != cur:
            self.cv.configure(cursor=cur)

    def on_arm(self, on: bool) -> None:
        if hasattr(self, "cv"):
            self.cv.configure(cursor="crosshair" if on else "arrow")

    def _clear_drop(self) -> None:
        self.cv.delete("drop")

    def _hover(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> bool:
        self._clear_drop()
        rc = self._cell_at(x_root, y_root)
        if rc is None:
            return False
        h, w = self.cv.winfo_height(), self.cv.winfo_width()
        y, x = y_root - self.cv.winfo_rooty(), x_root - self.cv.winfo_rootx()
        if y < 30:
            self.cv.yview_scroll(-1, "units")
        elif y > h - 30:
            self.cv.yview_scroll(1, "units")
        if x > w - 30:
            self.cv.xview_scroll(1, "units")
        elif x < 30:
            self.cv.xview_scroll(-1, "units")
        text = self._cell_text(rc)
        busy = bool(text) and not tt.PLACEHOLDER_RE.fullmatch(text)
        x1, y1, x2, y2 = self._box(rc)
        color = col("warning") if busy else col("success")
        self.cv.create_rectangle(x1 + 1, y1 + 1, x2 - 1, y2 - 1, outline=color, width=3, tags=("drop",))
        self.cv.create_text(x2 - 4, y1 + 2, anchor="ne", text=self._ref(rc), fill=color, font=font("caption"),
                            tags=("drop",))
        return True

    def _drop(self, x_root: int, y_root: int, payload: Dict[str, Any]) -> None:
        rc = self._cell_at(x_root, y_root)
        if rc is None:
            self.win.toast(t("Rilascia il campo su una cella del foglio."), "warning")
            return
        ref = self._ref(rc)
        other = self._key_at_cell(rc)
        if payload.get("kind") == "placed" and other == payload.get("key"):
            return
        text = self._cell_text(rc)
        if other is None and text and not tt.PLACEHOLDER_RE.fullmatch(text):
            if not messagebox.askyesno(
                    t("Cella con del testo"),
                    t("La cella {c} contiene «{v}».\nNel documento compilato verrà sostituita dal valore del "
                      "campo. Continuare?\n\n(Di solito il campo va nella cella accanto all'etichetta.)",
                      c=ref, v=_short(text, 40)), parent=self):
                return
        self.snapshot()
        if payload.get("kind") == "new":
            key = self.create_field(payload["type"], self._suggest(rc))
            if key is None:
                self.drop_snapshot()
                return
        else:
            key = payload["key"]
        if other is not None and other != key:
            self.mapping.pop(other, None)
        spec = self.schema["properties"].get(key, {})
        entry: Dict[str, Any] = {"type": "joined_cell" if tt.kind_of(spec) == "list" else "cell",
                                 "sheet": self.sheet, "cell": ref}
        if entry["type"] == "joined_cell":
            entry["separator"] = "\n"
        self.mapping[key] = entry
        self.selected = key
        self.changed()
        self._hide_welcome()

    # ------------------------------------------------------------ salvataggio
    def save(self) -> None:
        if not self.dirty:
            self.close()
            return
        props = self.schema.get("properties", {})
        mapping = {k: v for k, v in self.mapping.items() if k in props}
        for k, v in mapping.items():  # un campo diventato "Elenco" va in una cella unica, una riga per voce
            if isinstance(v, dict):
                if tt.kind_of(props[k]) == "list" and v.get("type") == "cell":
                    v.update({"type": "joined_cell", "separator": "\n"})
                elif tt.kind_of(props[k]) != "list" and v.get("type") == "joined_cell":
                    v["type"] = "cell"
                    v.pop("separator", None)
        self.win.mm.bump_version(self.mod.slug, new_schema=self._clean_schema(), new_mapping=mapping)
        self.win.refresh_modules(quiet=True)
        self.win.toast(t("Modulo salvato."), "success")
        self.close()
