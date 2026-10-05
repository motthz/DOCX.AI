"""DocumentGenerator: feature "Creazione documento".

Workflow (tutto sotto review gate):
1. Utente seleziona una cartella di riferimento (opzionale template DOCX)
2. DocumentLoader scansione ricorsiva → DocumentIndexer indicizza
3. DocumentRetriever seleziona brani pertinenti in base alla descrizione
4. AIService chiama LLM per generare bozza
5. UI mostra anteprima + fonti usate → conferma utente
6. Salvataggio su disco: DOCX + PDF opzionali + JSON bozza
"""

from __future__ import annotations

import json
import logging
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

from ..config import Config
from ..exporters.pdf_exporter import export_pdf
from ..module_manager import LoadedModule
from .ai_service import AIResult, AIService
from .document_indexer import Chunk, DocumentIndexer
from .document_loader import DocumentLoader, LoadedDocument
from .document_retriever import DocumentRetriever
from .source_tracker import SourceRef


LOG = logging.getLogger(__name__)


@dataclass
class GeneratedDocument:
    """Risultato generato (prima della conferma utente)."""

    title: str
    body_markdown: str
    sections: List[Dict[str, Any]] = field(default_factory=list)
    sources: List[SourceRef] = field(default_factory=list)
    placeholders_found: List[str] = field(default_factory=list)  # [DA DEFINIRE]
    template_used: Optional[Path] = None
    raw_llm_text: str = ""

    def sections_list(self) -> List[str]:
        return [s.get("title", "") for s in self.sections if s.get("title")]


class DocumentGenerator:
    OUTPUT_SCHEMA = {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "titolo": {"type": "string"},
            "sezioni": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "titolo": {"type": "string"},
                        "contenuto": {"type": "string"},
                        "sottosezioni": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "additionalProperties": False,
                                "properties": {
                                    "titolo": {"type": "string"},
                                    "contenuto": {"type": "string"},
                                },
                            },
                        },
                    },
                },
            },
        },
        "required": ["titolo", "sezioni"],
    }

    def __init__(
        self,
        config: Config,
        loader: DocumentLoader,
        indexer: DocumentIndexer,
        retriever: DocumentRetriever,
        ai: AIService,
    ):
        self.config = config
        self.loader = loader
        self.indexer = indexer
        self.retriever = retriever
        self.ai = ai

    # --------------------------------------------------------------
    # Public: genera bozza
    # --------------------------------------------------------------
    def generate(
        self,
        user_description: str,
        reference_folder: Optional[Path] = None,
        reference_files: Optional[List[Path]] = None,
        template_docx: Optional[Path] = None,
        module: Optional[LoadedModule] = None,
        *,
        progress_cb: Optional[Callable[[str], None]] = None,
        use_mock: bool = False,
    ) -> Tuple[GeneratedDocument, List[LoadedDocument]]:
        step = lambda msg: progress_cb(msg) if progress_cb else None

        # 1) Load references
        step("Analisi documenti di riferimento…")
        docs: List[LoadedDocument] = []
        if reference_folder is not None:
            docs.extend(self.loader.scan_folder(
                Path(reference_folder), recursive=True,
                progress_cb=lambda i, n, p: step(f"Caricamento file {i}/{n}: {p.name}"),
            ))
        if reference_files:
            for p in reference_files:
                try:
                    docs.append(self.loader.load_file(Path(p)))
                except Exception as exc:  # noqa: BLE001
                    LOG.warning("Scarto %s: %s", p, exc)

        # 2) Index
        step("Indicizzazione contenuti…")
        all_chunks: List[Chunk] = []
        for doc in docs:
            try:
                all_chunks.extend(self.indexer.index_document(doc))
            except Exception:  # noqa: BLE001
                pass

        # 3) Retrieve
        step("Ricerca informazioni pertinenti…")
        # Extra tokens: schema field names if module given
        extra_tokens: List[str] = []
        if module is not None:
            for k in (module.schema or {}).get("properties", {}).keys():
                extra_tokens.append(k)
        retrieved = self.retriever.retrieve(
            all_chunks, user_description,
            extra_query_tokens=extra_tokens, top_k=18, max_total_chars=12000,
        )

        # 4) AI call
        step("Generazione bozza tramite AI…")
        extra = (
            "Formatta l'output secondo lo schema indicato. "
            "Non inventare codici, responsabili, versioni o numeri. "
            "Ogni qual volta un'informazione non è disponibile nei documenti, "
            "scrivi letteralmente [DA DEFINIRE] all'interno del campo corrispondente. "
            "Cita tra parentesi la fonte (nome file) da cui hai ripreso ogni sezione o frase "
            "importante (es. (Fonte: procedura_01.pdf))."
        )
        result: AIResult = self.ai.extract_structured(
            feature_key="document_creation",
            schema=self.OUTPUT_SCHEMA,
            user_request=user_description,
            module=module,
            retrieved_chunks=retrieved,
            reference_docs=None,
            history_snippets=None,
            extra_headers=None,
            use_mock=use_mock,
        )
        step("Preparazione risultato…")
        if not result.success or not result.data:
            # Fallback: usa generate_free_text per avere qualcosa in mano
            ft = self.ai.generate_free_text(
                "document_creation", user_description,
                output_schema_hints=None,
                module=module,
                retrieved_chunks=retrieved,
                extra_instructions=(
                    "Produci un documento aziendale strutturato. "
                    "Usa titoli preceduti da ##, paragrafi, e tabelle quando utile. "
                    "Indica [DA DEFINIRE] per le informazioni mancanti. "
                    "Cita la fonte tra parentesi. " + extra
                ),
                max_tokens=3500,
                use_mock=use_mock,
            )
            title = Path(reference_folder).name if reference_folder else "Documento"
            body = ft.text_output or f"Errore durante la generazione: {result.error}"
            return (
                GeneratedDocument(
                    title=title, body_markdown=body, sections=[],
                    sources=ft.sources_used, raw_llm_text=ft.text_output,
                    template_used=template_docx,
                ),
                docs,
            )

        data = result.data or {}
        title = data.get("titolo") or "Documento"
        sezioni = data.get("sezioni") or []
        sections_list: List[Dict[str, Any]] = []
        body_lines: List[str] = [f"# {title}", ""]
        placeholders: List[str] = []
        for s in sezioni:
            if not isinstance(s, dict):
                continue
            t = s.get("titolo") or "Sezione"
            c = str(s.get("contenuto") or "")
            sections_list.append({"title": t, "content": c})
            body_lines.append(f"## {t}")
            body_lines.append("")
            body_lines.append(c)
            if "[DA DEFINIRE]" in c:
                placeholders.append(t)
            subs = s.get("sottosezioni") or []
            for ss in subs:
                if not isinstance(ss, dict):
                    continue
                st = ss.get("titolo") or ""
                sc = str(ss.get("contenuto") or "")
                if st:
                    body_lines.append(f"### {st}")
                    body_lines.append("")
                body_lines.append(sc)
                if "[DA DEFINIRE]" in sc:
                    placeholders.append(st or t)
            body_lines.append("")
        body = "\n".join(body_lines)
        return (
            GeneratedDocument(
                title=title,
                body_markdown=body,
                sections=sections_list,
                sources=result.sources_used,
                placeholders_found=sorted(set(placeholders)),
                template_used=template_docx,
                raw_llm_text=result.raw_text,
            ),
            docs,
        )

    # --------------------------------------------------------------
    # Public: salva dopo conferma
    # --------------------------------------------------------------
    def save_approved(
        self,
        generated: GeneratedDocument,
        output_dir: Path,
        *,
        make_docx: bool = True,
        make_pdf: bool = True,
    ) -> Dict[str, Path]:
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        stamp = time.strftime("%Y%m%d_%H%M%S")
        base_name = f"DocCreato_{stamp}"
        result_paths: Dict[str, Path] = {}

        # JSON
        json_path = output_dir / f"{base_name}.json"
        json_path.write_text(
            json.dumps(
                {
                    "title": generated.title,
                    "body_markdown": generated.body_markdown,
                    "sections": generated.sections,
                    "placeholders_found": generated.placeholders_found,
                    "sources": [
                        {
                            "file": s.file_name,
                            "page": s.page,
                            "sheet": s.sheet,
                            "cell": s.cell,
                            "excerpt": s.excerpt[:400],
                        }
                        for s in generated.sources
                    ],
                    "created_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                },
                indent=2, ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        result_paths["json"] = json_path

        # TXT (universale)
        txt_path = output_dir / f"{base_name}.txt"
        txt_path.write_text(generated.body_markdown, encoding="utf-8")
        result_paths["txt"] = txt_path

        # DOCX: crea un nuovo documento con sezioni
        if make_docx:
            try:
                import docx  # type: ignore
                out_docx = output_dir / f"{base_name}.docx"
                # Se template fornito, copiarlo; altrimenti nuovo documento vuoto
                if generated.template_used and Path(generated.template_used).is_file():
                    import shutil
                    shutil.copyfile(generated.template_used, out_docx)
                    d = docx.Document(str(out_docx))
                    from ..parsers.docx_parser import _PLACEHOLDER_RE, _iter_all_paragraphs, apply_placeholders
                    keys = {m.group(1) for p in _iter_all_paragraphs(d) for m in _PLACEHOLDER_RE.finditer(p.text)}
                    apply_placeholders(d, {
                        "titolo": generated.title,
                        "contenuto": generated.body_markdown,
                        "title": generated.title,
                        "body": generated.body_markdown,
                    })
                    # senza un segnaposto per il contenuto, il testo va in coda al modello
                    # (prima veniva aggiunto SEMPRE, quindi compariva due volte)
                    if not keys & {"contenuto", "body"}:
                        d.add_paragraph()
                        _markdown_to_docx(d, generated.body_markdown)
                    d.save(str(out_docx))
                else:
                    d = docx.Document()
                    _markdown_to_docx(d, generated.body_markdown)
                    d.save(str(out_docx))
                result_paths["docx"] = out_docx
            except Exception as exc:  # noqa: BLE001
                LOG.error("Salvataggio DOCX fallito: %s", exc)

        # PDF via ReportLab
        if make_pdf:
            try:
                pdf_path = output_dir / f"{base_name}.pdf"
                schema = {
                    "type": "object",
                    "properties": {
                        "Titolo": {"type": "string"},
                        "Sezioni": {"type": "string"},
                    },
                }
                data = {
                    "Titolo": generated.title,
                    "Sezioni": generated.body_markdown,
                }
                export_pdf(
                    pdf_path, data, schema,
                    report_id=base_name, module_name="Creazione Documento",
                )
                result_paths["pdf"] = pdf_path
            except Exception as exc:  # noqa: BLE001
                LOG.error("Salvataggio PDF fallito: %s", exc)

        return result_paths


def _markdown_to_docx(d: Any, markdown: str) -> None:
    """Scrive il testo della bozza (titoli #, ##, ###, elenchi, paragrafi) nel documento Word.
    Usa il testo completo e non solo le sezioni principali: le sottosezioni e la bozza
    di riserva in testo libero non vanno perse."""
    for raw in (markdown or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.lstrip()
        level = len(stripped) - len(stripped.lstrip("#"))
        if 1 <= level <= 4 and stripped[level:level + 1] == " ":
            d.add_heading(stripped[level:].strip(), level=level)
        elif stripped[:2] in ("- ", "* ", "• "):
            try:
                d.add_paragraph(stripped[2:].strip(), style="List Bullet")
            except KeyError:
                d.add_paragraph("• " + stripped[2:].strip())
        else:
            d.add_paragraph(stripped)
