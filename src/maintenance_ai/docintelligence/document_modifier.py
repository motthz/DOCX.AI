"""DocumentModifier: feature "Modifica documento".

Funzionalità:
- Import PDF (digitale o scannerizzato), DOCX, immagine, TXT
- Modalità A: richiesta testuale (es. "cambia frequenza")
- Modalità B: selezione testuale ("zona evidenziata") + istruzione associata
- Genera una NUOVA versione del documento (originale inalterato)
- Mostra confronto semplice prima/dopo
"""

from __future__ import annotations

import json
import logging
import shutil
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

from ..config import Config
from ..exporters.pdf_exporter import export_pdf
from .ai_service import AIResult, AIService
from .document_loader import DocumentLoader, LoadedDocument
from .source_tracker import SourceRef


LOG = logging.getLogger(__name__)


@dataclass
class TextSelection:
    """Selezione dell'utente: evidenziazione o zona indicizzata per paragrafo/pagina."""

    page: Optional[int] = None
    paragraph_index: Optional[int] = None
    substring: str = ""            # la sottostringa esatta richiesta
    user_label: str = ""           # label mostrata all'utente


@dataclass
class ModificationRequest:
    """Una singola modifica richiesta."""

    instruction: str
    selection: Optional[TextSelection] = None

    def describe(self) -> str:
        if self.selection is None:
            return self.instruction
        sel = self.selection
        parts = []
        if sel.page is not None:
            parts.append(f"pagina {sel.page+1}")
        if sel.substring:
            parts.append(f"riferimento a '{sel.substring[:80]}'")
        if not parts:
            return self.instruction
        return self.instruction + " (" + ", ".join(parts) + ")"


@dataclass
class ModifiedDocument:
    """Risultato modifica (bozza prima di conferma)."""

    original_text: str
    modified_text: str
    changes_summary: List[str] = field(default_factory=list)
    sources: List[SourceRef] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    raw_llm_text: str = ""
    format_preserved: bool = True

    def diff_lines(self) -> List[Tuple[str, str]]:
        """Semplice line-by-line diff per UI (added/removed lines)."""
        import difflib
        orig = (self.original_text or "").splitlines()
        modif = (self.modified_text or "").splitlines()
        out: List[Tuple[str, str]] = []
        for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, orig, modif).get_opcodes():
            if tag == "equal":
                continue
            for line in orig[i1:i2]:
                out.append(("-", line))
            for line in modif[j1:j2]:
                out.append(("+", line))
        return out


class DocumentModifier:
    MODIFY_SCHEMA = {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "testo_modificato": {"type": "string"},
            "modifiche": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "zona": {"type": "string"},
                        "prima": {"type": "string"},
                        "dopo": {"type": "string"},
                        "motivazione": {"type": "string"},
                    },
                },
            },
            "avvertenze": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["testo_modificato", "modifiche", "avvertenze"],
    }

    def __init__(
        self,
        config: Config,
        loader: DocumentLoader,
        ai: AIService,
    ):
        self.config = config
        self.loader = loader
        self.ai = ai

    # --------------------------------------------------------------
    # Public: carica + modifica
    # --------------------------------------------------------------
    def load_original(self, path: Path) -> LoadedDocument:
        return self.loader.load_file(path)

    def modify(
        self,
        original_doc: LoadedDocument,
        requests: List[ModificationRequest],
        *,
        progress_cb: Optional[Callable[[str], None]] = None,
        module_folder: Optional[Path] = None,
        use_mock: bool = False,
    ) -> ModifiedDocument:
        step = lambda msg: progress_cb(msg) if progress_cb else None
        step("Analisi testo originale…")
        orig_text = original_doc.full_text or ""
        if not orig_text.strip():
            raise RuntimeError(
                "Il documento caricato non contiene testo estraibile. "
                "Verificare che non sia una scansione senza OCR disponibile."
            )

        step("Preparazione istruzioni modifica…")
        # Costruisci la richiesta
        user_lines: List[str] = []
        user_lines.append("=== TESTO ORIGINALE DEL DOCUMENTO ===")
        # Tronca se troppo lungo
        from ..llm.prompt_builder import _truncate
        orig_show = _truncate(orig_text, 18000, "originale")
        user_lines.append(orig_show)
        user_lines.append("=== RICHIESTE DI MODIFICA ===")
        for i, req in enumerate(requests, 1):
            user_lines.append(f"{i}. {req.describe()}")
        user_lines.append(
            "Regole: Modifica SOLAMENTE le parti richieste e le parti "
            "strettamente necessarie per mantenere la coerenza. "
            "Non riscrivere arbitrariamente parti non toccate. "
            "Se la formattazione non può essere preservata perfettamente, "
            "aggiungi un avvertimento nella lista 'avvertenze'. "
            "Restituisci lo schema indicato."
        )
        step("Applicazione modifiche tramite AI…")
        result: AIResult = self.ai.extract_structured(
            feature_key="document_edit",
            schema=self.MODIFY_SCHEMA,
            user_request="\n\n".join(user_lines),
            module=None,
            retrieved_chunks=None,
            reference_docs=None,
            history_snippets=None,
            use_mock=use_mock,
        )
        step("Preparazione confronto…")
        warnings: List[str] = []
        changes: List[str] = []
        modified_text = orig_text
        if not result.success or not result.data:
            # Fallback: genera free text
            ft = self.ai.generate_free_text(
                "document_edit", "\n\n".join(user_lines),
                output_schema_hints=None,
                extra_instructions=(
                    "Applica le modifiche richieste. Restituisci nella tua risposta "
                    "il testo completo modificato del documento, nella sua interezza. "
                    "Evidenzia le modifiche con [[prima → dopo]] tra parentesi quadre "
                    "all'interno del testo, poi alla fine scrivi '--- MODIFICHE ---' "
                    "e un elenco riassuntivo."
                ),
                max_tokens=6000,
                use_mock=use_mock,
            )
            raw = ft.text_output or ""
            # Extract: separa il testo modificato dalle modifiche
            parts = raw.split("--- MODIFICHE ---", 1)
            modified_text = parts[0].strip() or orig_text
            if len(parts) > 1:
                changes.append(parts[1].strip()[:500])
            warnings.append("Formattazione modificata: elaborazione fallback free-text.")
            return ModifiedDocument(
                original_text=orig_text,
                modified_text=modified_text,
                changes_summary=changes or ["Modifiche applicate in modalità testo libero."],
                warnings=warnings,
                sources=ft.sources_used,
                raw_llm_text=raw,
                format_preserved=False,
            )
        data = result.data or {}
        modified_text = str(data.get("testo_modificato") or orig_text)
        modifiche = data.get("modifiche") or []
        for m in modifiche:
            if not isinstance(m, dict):
                continue
            try:
                zona = m.get("zona", "")
                prima = str(m.get("prima", ""))[:200]
                dopo = str(m.get("dopo", ""))[:200]
                changes.append(f"{zona}: «{prima}» → «{dopo}»")
            except Exception:  # noqa: BLE001
                pass
        avvertenze = data.get("avvertenze") or []
        for a in avvertenze:
            if a:
                warnings.append(str(a))
        return ModifiedDocument(
            original_text=orig_text,
            modified_text=modified_text,
            changes_summary=changes,
            warnings=warnings,
            sources=result.sources_used,
            raw_llm_text=result.raw_text,
            format_preserved=not any("formattaz" in w.lower() for w in warnings),
        )

    # --------------------------------------------------------------
    # Public: salva versione nuova (originale non toccato)
    # --------------------------------------------------------------
    def save_new_version(
        self,
        original_path: Path,
        modified: ModifiedDocument,
        output_dir: Path,
    ) -> Dict[str, Path]:
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        stamp = time.strftime("%Y%m%d_%H%M%S")
        orig = Path(original_path)
        base = f"{orig.stem}_v2_{stamp}"
        result: Dict[str, Path] = {}

        # 1) TXT new version
        txt_path = output_dir / f"{base}.txt"
        txt_path.write_text(modified.modified_text, encoding="utf-8")
        result["txt"] = txt_path

        # 2) JSON con metadata e diff
        meta_path = output_dir / f"{base}.json"
        meta_path.write_text(
            json.dumps({
                "original": str(original_path),
                "version": "2.0",
                "created_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
                "changes_summary": modified.changes_summary,
                "warnings": modified.warnings,
                "modified_text": modified.modified_text,
            }, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        result["json"] = meta_path

        # 3) Se l'originale era DOCX, prova a creare un nuovo DOCX
        if orig.suffix.lower() == ".docx":
            try:
                import docx  # type: ignore
                out_docx = output_dir / f"{base}.docx"
                shutil.copyfile(orig, out_docx)
                # Try to replace content paragraphs one-by-one (best effort)
                try:
                    d = docx.Document(str(out_docx))
                    lines = modified.modified_text.splitlines()
                    # Sostituisci i primi N paragrafi con le linee corrispondenti
                    for i, p in enumerate(d.paragraphs):
                        if i < len(lines):
                            # Preserva gli stili, sostituisci il testo
                            if p.runs:
                                first_run = p.runs[0]
                                for r in p.runs[1:]:
                                    r.text = ""
                                first_run.text = lines[i]
                            else:
                                p.text = lines[i]
                    d.save(str(out_docx))
                    result["docx"] = out_docx
                except Exception as exc:  # noqa: BLE001
                    modified.warnings.append(f"Sostituzione paragrafi DOCX fallita: {exc}")
                    # Salva comunque il file copiato (intatto)
                    if not out_docx.is_file():
                        shutil.copyfile(orig, out_docx)
                        result["docx"] = out_docx
            except Exception as exc:  # noqa: BLE001
                LOG.warning("Salvataggio nuova versione DOCX fallito: %s", exc)

        # 4) PDF reportabile del nuovo testo
        try:
            pdf_path = output_dir / f"{base}.pdf"
            schema = {
                "type": "object",
                "properties": {
                    "Titolo": {"type": "string"},
                    "Testo": {"type": "string"},
                },
            }
            export_pdf(
                pdf_path,
                {"Titolo": f"{orig.name} - Versione modificata",
                 "Testo": modified.modified_text},
                schema,
                report_id=base,
                module_name="Modifica Documento",
            )
            result["pdf"] = pdf_path
        except Exception as exc:  # noqa: BLE001
            LOG.warning("Salvataggio PDF modifica fallito: %s", exc)

        return result
