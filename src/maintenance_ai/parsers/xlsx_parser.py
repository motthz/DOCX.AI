"""XLSX parser / text extractor / cell writer.

Uses openpyxl (read_only=False so we can also write). Enforces defusedxml if
available, and never evaluates formulas.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List

try:
    from defusedxml.common import DTDForbidden, EntitiesForbidden
except Exception:  # pragma: no cover
    DTDForbidden = Exception  # type: ignore[assignment,misc]
    EntitiesForbidden = Exception  # type: ignore[assignment,misc]

from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class XlsxExtraction:
    full_text: str
    sheets: Dict[str, List[List[Any]]] = field(default_factory=dict)
    placeholders: List[str] = field(default_factory=list)
    meta: Dict[str, Any] = field(default_factory=dict)


def load_workbook_safe(path: Path) -> Workbook:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"XLSX non trovato: {path}")
    # data_only = True -> le formule restituiscono l'ultimo valore calcolato,
    # non vengono valutate da openpyxl e non vengono eseguite.
    return load_workbook(str(path), data_only=True, keep_vba=False, read_only=False)


def extract_text(wb: Workbook) -> XlsxExtraction:
    out = XlsxExtraction(full_text="")
    import re
    ph_re = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_\.]*)\s*\}\}")
    placeholders = set()
    pieces: List[str] = []
    for ws in wb.worksheets:
        title = ws.title or ""
        pieces.append(f"[Sheet: {title}]")
        rows: List[List[Any]] = []
        for row in ws.iter_rows(values_only=True):
            row_values: List[Any] = list(row)
            rows.append(row_values)
            cell_strs = []
            for value in row_values:
                s = "" if value is None else str(value)
                cell_strs.append(s)
                placeholders.update(ph_re.findall(s))
            pieces.append(" | ".join(cell_strs))
        out.sheets[title] = rows
    out.full_text = "\n".join(pieces)
    out.placeholders = sorted(placeholders)
    out.meta["sheet_count"] = len(wb.worksheets)
    return out


def _cell_value_for_write(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(_line_for(v) for v in value)
    if isinstance(value, dict):
        return "\n".join(f"{k}: {_line_for(v)}" for k, v in value.items())
    return value


def _line_for(value: Any) -> str:
    if isinstance(value, dict):
        return " ; ".join(f"{k}={v}" for k, v in value.items())
    return "" if value is None else str(value)


def apply_mapping(
    wb: Workbook,
    mapping: Dict[str, Any],
    values: Dict[str, Any],
) -> List[str]:
    """Apply a mapping file of the form::

        {"field": {"type": "cell", "sheet": "Rapporto", "cell": "B4"},
         "field2": {"type": "joined_cell", "sheet": "Rapporto", "cell": "B11", "separator": "\\n"}}

    Returns a list of applied field names.
    """
    applied: List[str] = []
    sheets: Dict[str, Worksheet] = {ws.title: ws for ws in wb.worksheets}
    for fld, spec in mapping.items():
        if fld not in values or not isinstance(spec, dict):
            continue
        sheet_name = spec.get("sheet")
        cell_ref = spec.get("cell") or ""
        if not cell_ref:
            continue
        # FR51: multi-sheet SheetName!Cell bang-split backward compat
        if "!" in cell_ref and not sheet_name:
            split = cell_ref.split("!", 1)
            sheet_name = split[0].strip()
            cell_ref = split[1].strip()
        sheet = sheets.get(sheet_name) if sheet_name else wb.active
        if sheet is None:
            continue
        typ = spec.get("type", "cell")
        if not cell_ref:
            continue
        value = values[fld]
        if typ == "joined_cell":
            sep = spec.get("separator", "\n")
            if isinstance(value, list):
                text = sep.join(_line_for(v) for v in value)
            else:
                text = _line_for(value)
            sheet[cell_ref] = text
        else:
            sheet[cell_ref] = _cell_value_for_write(value)
        applied.append(fld)
    return applied
