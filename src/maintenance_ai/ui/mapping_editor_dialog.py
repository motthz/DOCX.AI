"""Mapping editor dialog for XLSX templates.

Shows the mapping as a Treeview (Field / Type / Sheet / Cell) and a small
XLSX sheet/cell preview. Double-click a cell in the sheet preview to assign
it to the currently selected mapping row.
"""

from __future__ import annotations

import json
import tkinter as tk
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from tkinter import ttk
except Exception:  # noqa: BLE001
    ttk = None  # type: ignore


class MappingEditorDialog:
    """Modal dialog. Call run() → (accepted: bool, new_mapping_dict)."""

    def __init__(self, master, template_path: Path, mapping: Dict[str, Any],
                 schema: Optional[Dict[str, Any]] = None,
                 title: str = "Modifica mapping XLSX"):
        self.template_path = Path(template_path)
        self.result = False
        self.new_mapping: Dict[str, Any] = {}
        self._mapping = {k: dict(v) if isinstance(v, dict) else v for k, v in (mapping or {}).items()}
        self._schema = schema or {}
        self.top = tk.Toplevel(master)
        self.top.title(title)
        self.top.geometry("1020x620")
        self.top.transient(master)
        self.top.grab_set()
        self._build()

    def _build(self) -> None:
        # Left: mapping treeview
        pad = {"padx": 10, "pady": 6}
        left = ttk.LabelFrame(self.top, text="Mappatura campi → celle")
        left.pack(side="left", fill="both", expand=True, **pad)
        cols = ("field", "type", "sheet", "cell", "sep")
        tv = ttk.Treeview(left, columns=cols, show="headings")
        for c, label, w in (("field", "Campo", 140), ("type", "Tipo", 90),
                            ("sheet", "Foglio", 120), ("cell", "Cella", 90),
                            ("sep", "Sep.", 80)):
            tv.heading(c, text=label)
            tv.column(c, width=w, anchor="w")
        vsb = ttk.Scrollbar(left, orient="vertical", command=tv.yview)
        tv.configure(yscrollcommand=vsb.set)
        tv.pack(side="left", fill="both", expand=True, padx=(8, 0), pady=8)
        vsb.pack(side="right", fill="y", pady=8)
        self._tv = tv

        # Right: XLSX preview
        right = ttk.LabelFrame(self.top, text=f"Template XLSX: {self.template_path.name} — doppio click cella per assegnare")
        right.pack(side="right", fill="both", expand=True, **pad)
        top_row = ttk.Frame(right)
        top_row.pack(fill="x", padx=8, pady=(6, 4))
        ttk.Label(top_row, text="Foglio:").pack(side="left")
        self._sheet_var = tk.StringVar()
        self._sheet_cb = ttk.Combobox(top_row, textvariable=self._sheet_var, state="readonly", width=30)
        self._sheet_cb.pack(side="left", padx=6)
        self._sheet_cb.bind("<<ComboboxSelected>>", self._reload_sheet)
        tv2 = ttk.Treeview(right, show="headings", height=22)
        sb = ttk.Scrollbar(right, orient="vertical", command=tv2.yview)
        hsb = ttk.Scrollbar(right, orient="horizontal", command=tv2.xview)
        tv2.configure(yscrollcommand=sb.set, xscrollcommand=hsb.set)
        self._preview_tv = tv2
        self._preview_tv.bind("<Double-1>", self._on_cell_dblclick)
        tv2.grid(row=1, column=0, sticky="nsew", padx=(8, 0), pady=(0, 8))
        sb.grid(row=1, column=1, sticky="ns", pady=(0, 8))
        hsb.grid(row=2, column=0, columnspan=2, sticky="ew", padx=(8, 0))
        right.rowconfigure(1, weight=1)
        right.columnconfigure(0, weight=1)

        # Buttons bottom
        bb = ttk.Frame(self.top)
        bb.pack(side="bottom", fill="x", padx=10, pady=8)
        ttk.Button(bb, text="Agg. da schema", command=self._auto_from_schema).pack(side="left", padx=4)
        ttk.Button(bb, text="Elimina selez.", command=self._delete_selected).pack(side="left", padx=4)
        ttk.Button(bb, text="Annulla", command=self._on_cancel).pack(side="right", padx=4)
        ttk.Button(bb, text="Salva", style="Accent.TButton", command=self._on_save).pack(side="right", padx=4)
        self._load_mapping_rows()
        self._load_workbook()

    # --------------------------------------------------------------
    def _auto_from_schema(self) -> None:
        props = (self._schema or {}).get("properties", {}) or {}
        existing = set(self._tv.set(r, "field") for r in self._tv.get_children(""))
        for k in props.keys():
            if k not in existing:
                self._tv.insert("", "end", values=(k, "cell", "", "", ""))

    def _delete_selected(self) -> None:
        for iid in self._tv.selection():
            self._tv.delete(iid)

    def _load_mapping_rows(self) -> None:
        for r in self._tv.get_children(""):
            self._tv.delete(r)
        for field, spec in self._mapping.items():
            if not isinstance(spec, dict):
                continue
            typ = str(spec.get("type", "cell"))
            sheet = str(spec.get("sheet", ""))
            cell = str(spec.get("cell", ""))
            sep = str(spec.get("separator", "")) if typ == "joined_cell" else ""
            self._tv.insert("", "end", values=(field, typ, sheet, cell, sep))

    # --------------------------------------------------------------
    def _load_workbook(self) -> None:
        try:
            from ..parsers.xlsx_parser import load_workbook_safe
            wb = load_workbook_safe(self.template_path)
        except Exception:  # noqa: BLE001
            self._wb = None
            return
        self._wb = wb
        sheets = list(wb.sheetnames)
        self._sheet_cb["values"] = sheets
        if sheets:
            self._sheet_var.set(sheets[0])
            self._reload_sheet()

    def _reload_sheet(self, _evt=None) -> None:
        for r in self._preview_tv.get_children(""):
            self._preview_tv.delete(r)
        for col in self._preview_tv["columns"]:
            self._preview_tv.heading(col, text="")
            self._preview_tv.column(col, width=0)
        if self._wb is None:
            return
        name = self._sheet_var.get()
        if name not in self._wb.sheetnames:
            return
        ws = self._wb[name]
        # Limit rows/cols for preview performance
        max_col = min(ws.max_column, 26)
        max_row = min(ws.max_row, 200)
        cols_list = [f"C{i}" for i in range(1, max_col + 1)]
        self._preview_tv["columns"] = cols_list
        from openpyxl.utils import get_column_letter
        for i, c in enumerate(cols_list, 1):
            letter = get_column_letter(i)
            self._preview_tv.heading(c, text=letter)
            self._preview_tv.column(c, width=90, anchor="w", stretch=False)
        for r_idx in range(1, max_row + 1):
            values = []
            for c_idx in range(1, max_col + 1):
                v = ws.cell(row=r_idx, column=c_idx).value
                values.append("" if v is None else str(v))
            self._preview_tv.insert("", "end", iid=f"R{r_idx}", values=values)
        self._current_sheet = name

    def _on_cell_dblclick(self, evt=None) -> None:
        sel = self._tv.selection()
        if not sel:
            tk.messagebox.showinfo("Seleziona mapping",
                                   "Selezionare una riga nella mappatura a sinistra,\n"
                                   "poi doppio click sulla cella corrispondente qui a destra.",
                                   parent=self.top)
            return
        iid = sel[0]
        field, typ, _s, _c, _sep = self._tv.item(iid, "values")
        try:
            preview_item = self._preview_tv.identify_row(evt.y)
            col_idx = int(self._preview_tv.identify_column(evt.x).replace("#", ""))
        except Exception:  # noqa: BLE001
            return
        if not preview_item or not preview_item.startswith("R"):
            return
        try:
            row_idx = int(preview_item[1:])
        except ValueError:
            return
        from openpyxl.utils import get_column_letter
        cell_ref = f"{get_column_letter(col_idx)}{row_idx}"
        sheet = getattr(self, "_current_sheet", self._sheet_var.get())
        # Update existing row
        self._tv.item(iid, values=(field, typ, sheet, cell_ref, _sep))
        self._tv.see(iid)

    # --------------------------------------------------------------
    def _collect_mapping(self) -> Dict[str, Any]:
        mapping: Dict[str, Any] = {}
        for iid in self._tv.get_children(""):
            field, typ, sheet, cell, sep = self._tv.item(iid, "values")
            if not field or not cell:
                continue
            entry: Dict[str, Any] = {"type": str(typ) or "cell",
                                     "sheet": str(sheet) or "",
                                     "cell": str(cell)}
            if typ == "joined_cell" and sep:
                entry["separator"] = str(sep)
            mapping[str(field)] = entry
        return mapping

    def _on_save(self) -> None:
        self.new_mapping = self._collect_mapping()
        self.result = True
        self.top.destroy()

    def _on_cancel(self) -> None:
        self.result = False
        self.new_mapping = {}
        self.top.destroy()

    def run(self) -> tuple:
        self.top.wait_window()
        return (self.result, self.new_mapping)
