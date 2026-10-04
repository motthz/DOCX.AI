"""Export Excel riepilogativo di piu' rapporti (es. tutti quelli di un periodo).

Un foglio "Rapporti" con una riga per rapporto e una colonna per ogni campo
(unione dei campi di tutti i rapporti), piu' un foglio "Riepilogo" con i
conteggi per modulo e per stato.
"""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_IT = {"draft": "Bozza", "approved": "Approvato", "exported": "Esportato", "failed": "Fallito"}
HEADER_FILL = PatternFill("solid", fgColor="1E3A8A")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _cell(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        if isinstance(value, list) and all(isinstance(v, dict) for v in value):
            return "; ".join(", ".join(f"{k}: {v}" for k, v in item.items()) for item in value)
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "Sì" if value else "No"
    return value


def export_summary(rows: List[Dict[str, Any]], dest: Path, *, title: str = "") -> Path:
    data_rows: List[Dict[str, Any]] = []
    fields: List[str] = []
    for r in rows:
        try:
            data = json.loads(r.get("final_json") or r.get("draft_json") or "{}")
        except (TypeError, ValueError):
            data = {}
        if not isinstance(data, dict):
            data = {}
        for k in data:
            if k not in fields:
                fields.append(k)
        data_rows.append(data)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rapporti"
    base_cols = ["ID", "Modulo", "Stato", "Creato", "Approvato", "Descrizione intervento"]
    headers = base_cols + [f.replace("_", " ").capitalize() for f in fields]
    ws.append(headers)
    for r, data in zip(rows, data_rows):
        ws.append([r.get("id"), r.get("module"), STATUS_IT.get(r.get("status", ""), r.get("status")),
                   (r.get("created_at") or "").replace("T", " "), (r.get("approved_at") or "").replace("T", " "),
                   r.get("input_description")] + [_cell(data.get(f)) for f in fields])
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        lengths = [len(str(c.value or "")) for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i)]
        width = max([len(str(h)), 10] + lengths)
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)

    sm = wb.create_sheet("Riepilogo")
    sm.append([title or "Riepilogo rapporti"])
    sm["A1"].font = Font(bold=True, size=14)
    sm.append([])
    sm.append(["Totale rapporti", len(rows)])
    sm.append([])
    sm.append(["Modulo", "Rapporti"])
    for mod, n in Counter(r.get("module") or "-" for r in rows).most_common():
        sm.append([mod, n])
    sm.append([])
    sm.append(["Stato", "Rapporti"])
    for st, n in Counter(STATUS_IT.get(r.get("status", ""), r.get("status")) for r in rows).most_common():
        sm.append([st, n])
    sm.column_dimensions["A"].width = 40
    sm.column_dimensions["B"].width = 14
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
