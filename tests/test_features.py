"""Test delle funzioni v0.3: duplicazione, versioni, foto nel PDF, export Excel,
import CSV, backup/ripristino, pulizia, diagnostica, impostazioni, traduzioni."""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))


class AppTestCase(unittest.TestCase):
    """App completa su una cartella dati temporanea con il modulo di esempio."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="mai_feat_"))
        cls.env = mock.patch.dict(os.environ, {"DOCX_AI_DATA_DIR": str(cls.tmp)})
        cls.env.start()
        shutil.copytree(REPO / "examples" / "modules", cls.tmp / "workspace" / "modules", dirs_exist_ok=True)
        from docx_ai.app import App
        cls.app = App.bootstrap()
        cls.mod = cls.app.module_manager.list_modules()[0]

    @classmethod
    def tearDownClass(cls):
        cls.app.shutdown()
        cls.env.stop()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def _report(self, data=None, status="draft", desc="Sostituita cinghia compressore C-12") -> int:
        data = data or {"apparecchiatura": "C-12", "note": "ok"}
        return self.app.db.create_report(module_id=self.mod.id, module_version=self.mod.version, status=status,
                                         input_description=desc, draft_json=json.dumps(data))


class ReportFeatureTests(AppTestCase):
    def test_duplicate_report(self):
        rid = self._report()
        new = self.app.reports.duplicate_report(rid)
        row = self.app.db.get_report(new)
        self.assertEqual(row["status"], "draft")
        self.assertEqual(json.loads(row["draft_json"])["apparecchiatura"], "C-12")
        self.assertEqual(row["source"], "duplicate")

    def test_versions_and_restore(self):
        rid = self._report()
        self.app.reports.approve_draft(rid, {"apparecchiatura": "C-12", "note": "prima"})
        self.app.reports.update_approved(rid, {"apparecchiatura": "C-12", "note": "seconda"})
        versions = self.app.db.list_report_versions(rid)
        self.assertEqual([v["version"] for v in versions], [2, 1])
        self.app.reports.restore_version(rid, versions[-1])
        row = self.app.db.get_report(rid)
        self.assertEqual(json.loads(row["final_json"])["note"], "prima")
        self.assertEqual(len(self.app.db.list_report_versions(rid)), 3)

    def test_photos_end_up_in_pdf_and_export_folder(self):
        from PIL import Image
        img = self.tmp / "foto.png"
        Image.new("RGB", (800, 600), (200, 30, 30)).save(img)
        rid = self._report()
        self.assertEqual(self.app.reports.add_photos(rid, [img]), 1)
        self.app.reports.approve_draft(rid, {k: "x" for k in (self.mod.schema or {}).get("properties", {})})
        json_p, _doc, pdf = self.app.reports.finalize_exports(rid, self.mod)
        self.assertTrue(pdf.is_file())
        self.assertTrue(any((json_p.parent / "foto").glob("foto_01.*")))
        from pypdf import PdfReader
        reader = PdfReader(str(pdf))
        images = sum(len(p.images) for p in reader.pages)
        self.assertGreaterEqual(images, 1)

    def test_search_filters(self):
        rid = self._report({"firma_operatore_manutenzione": "Giulia Verdi", "reparto": "Verniciatura"},
                           desc="Controllo nastro trasportatore")
        rows, _n = self.app.db.search_reports(query="nastro")
        self.assertIn(rid, [r["id"] for r in rows])
        rows, _n = self.app.db.search_reports(field_filters={"tecnico": "verdi", "impianto": "vernic"})
        self.assertEqual([r["id"] for r in rows], [rid])


class TableTests(AppTestCase):
    def test_excel_summary(self):
        from openpyxl import load_workbook
        from docx_ai.exporters.summary_xlsx import export_summary
        self._report()
        rows, _n = self.app.db.search_reports(limit=1000)
        out = export_summary(rows, self.tmp / "riepilogo.xlsx", title="Test")
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Documenti", "Riepilogo"])
        self.assertEqual(wb["Documenti"].max_row, len(rows) + 1)

    def test_csv_import(self):
        from docx_ai.services import import_service
        csv_p = self.tmp / "interventi.csv"
        csv_p.write_text("Data;Apparecchiatura;Note;Descrizione\n02/10/2026;C-12;Cinghia;Cambio cinghia\n"
                         "03/10/2026;P-4;Filtro;Pulizia filtro\n", encoding="utf-8")
        table = import_service.read_table(csv_p)
        self.assertEqual(len(table.rows), 2)
        mapping = import_service.suggest_mapping(table.headers, self.mod.schema)
        self.assertEqual(mapping.get("apparecchiatura"), "Apparecchiatura")
        ids = import_service.import_rows(self.app.db, self.mod, table, mapping, status="approved",
                                         description_column="Descrizione")
        self.assertEqual(len(ids), 2)
        row = self.app.db.get_report(ids[0])
        self.assertEqual(row["status"], "approved")
        self.assertEqual(json.loads(row["final_json"])["apparecchiatura"], "C-12")
        self.assertEqual(row["input_description"], "Cambio cinghia")


class BackupTests(AppTestCase):
    def test_backup_restore_roundtrip(self):
        from docx_ai.services import backup_service
        rid = self._report(desc="Rapporto da salvare nel backup")
        zip_p = backup_service.create_backup(self.app.config, self.app.db, self.tmp / "b.zip")
        with zipfile.ZipFile(zip_p) as zf:
            names = zf.namelist()
        self.assertIn("docx_ai.db", names)
        self.assertTrue(any(n.startswith("workspace/") for n in names))
        info = backup_service.validate_backup(zip_p)
        self.assertEqual(info["app"], "DOCX.AI")
        # ripristino su una cartella dati separata
        target = self.tmp / "restored"
        target.mkdir()
        shutil.copy2(zip_p, target / backup_service.PENDING)
        msg = backup_service.apply_pending_restore(target)
        self.assertIn("ripristinato", msg)
        from docx_ai.db import Database
        db = Database(target / "docx_ai.db")
        try:
            self.assertIsNotNone(db.get_report(rid))
        finally:
            db.close()

    def test_invalid_backup_rejected(self):
        from docx_ai.services import backup_service
        bad = self.tmp / "bad.zip"
        with zipfile.ZipFile(bad, "w") as zf:
            zf.writestr("x.txt", "x")
        with self.assertRaises(ValueError):
            backup_service.validate_backup(bad)


class SupportTests(AppTestCase):
    def test_diagnostics_package_has_no_reports(self):
        from docx_ai.services import diagnostics
        self._report(desc="DATO RISERVATO DEL CLIENTE")
        p = diagnostics.create_package(self.app, self.tmp / "diag")
        with zipfile.ZipFile(p) as zf:
            names = zf.namelist()
            blob = b"".join(zf.read(n) for n in names)
        self.assertIn("sistema.json", names)
        self.assertNotIn(b"DATO RISERVATO", blob)
        self.assertFalse(any(n.endswith(".db") for n in names))
        url = diagnostics.issue_url(self.app)
        self.assertTrue(url.startswith("https://github.com/motthz/DOCX.AI/issues/new?"))

    def test_housekeeping(self):
        from docx_ai.services import housekeeping
        logs = self.app.config.logs_root()
        crash = logs / "crash.log"
        crash.write_bytes(b"x" * (housekeeping.CRASH_LOG_MAX + 1000))
        housekeeping.run(self.app.config.data_root)
        self.assertLess(crash.stat().st_size, housekeeping.CRASH_LOG_MAX)

    def test_settings_typed(self):
        s = self.app.settings
        self.assertEqual(s.get("ai.idle_minutes"), 15)
        s.set("ai.preload", False)
        self.assertIs(s.get("ai.preload"), False)
        s.set("ui.scale", 1.25)
        self.assertAlmostEqual(s.get("ui.scale"), 1.25)

    def test_db_maintenance(self):
        info = self.app.db.maintenance()
        self.assertTrue(info["optimized"])


class I18nTests(unittest.TestCase):
    def tearDown(self):
        from docx_ai.ui import i18n
        i18n.set_language("it-IT")

    def test_translation_and_format(self):
        from docx_ai.ui import i18n
        i18n.set_language("en-US")
        self.assertEqual(i18n.t("Salva bozza"), "Save draft")
        self.assertEqual(i18n.t("{n} bozze", n=3), "3 drafts")
        self.assertEqual(i18n.t("Testo non nel catalogo"), "Testo non nel catalogo")
        i18n.set_language("it-IT")
        self.assertEqual(i18n.t("Salva bozza"), "Salva bozza")

    def test_catalog_complete_and_consistent(self):
        import re
        cat = json.loads((REPO / "src" / "docx_ai" / "locales" / "en.json").read_text(encoding="utf-8"))
        self.assertTrue(all(cat.values()))
        for k, v in cat.items():
            self.assertEqual(sorted(re.findall(r"(?<!\{)\{\w+\}(?!\})", k)),
                             sorted(re.findall(r"(?<!\{)\{\w+\}(?!\})", v)), k)


class HardwareTests(unittest.TestCase):
    def test_detect(self):
        from docx_ai.llm import hardware
        hw = hardware.detect()
        self.assertGreater(hw.cpu_cores, 0)
        self.assertIn(hw.recommended_model(), ("Qwen3-1.7B-Q8_0.gguf", "Qwen3-0.6B-Q8_0.gguf"))
        self.assertIn("CPU", hw.summary())


if __name__ == "__main__":
    unittest.main()
