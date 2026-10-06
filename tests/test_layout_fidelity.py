"""Editor visuale fedele al modulo: posizioni dei paragrafi sulla pagina (Word) e
geometria/stili del foglio (Excel)."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))


def _sample_doc():
    import docx
    d = docx.Document()
    d.add_heading("Verbale di intervento", 1)
    d.add_paragraph("Cliente: ______")
    d.add_paragraph("")
    tb = d.add_table(rows=2, cols=2)
    tb.cell(0, 0).text = "Nome"
    tb.cell(1, 0).text = "Data"
    tb.cell(1, 1).text = "{{data}}"
    d.add_paragraph("Firma “tecnico”")
    return d


def _fake_render(copy_doc, pdf_path: Path, skip_empty: int = -1) -> None:
    """Impagina in modo semplice la copia (come farebbe Word): testo in ordine, tabella con bordi.
    ``skip_empty``: un paragrafo vuoto che il "motore" non disegna (es. cella unita)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from docx_ai import docx_layout as dl
    infos = dl.paragraphs(copy_doc)
    c = canvas.Canvas(str(pdf_path), pagesize=A4)
    _w, h = A4
    y = h - 72
    c.drawString(30, h - 30, "1")  # numero di pagina: testo in piu' nel PDF
    cells = [i for i in infos if i.in_cell]
    table_done = False
    for idx, info in enumerate(infos):
        if info.in_cell:
            if table_done:
                continue
            table_done = True
            top = y
            for r in range(2):
                for col in range(2):
                    x0, yy = 72 + col * 200, top - r * 20
                    c.line(x0, yy, x0 + 200, yy)
                    c.line(x0, yy - 20, x0 + 200, yy - 20)
                    c.line(x0, yy, x0, yy - 20)
                    c.line(x0 + 200, yy, x0 + 200, yy - 20)
            for k, ci in enumerate(cells):
                r, col = divmod(k, 2)
                text = ci.paragraph.text
                small = text.startswith(dl.MARK)
                c.setFont("Helvetica", 1 if small else 12)
                c.drawString(72 + col * 200 + 4, top - r * 20 - 14, text)
            c.setFont("Helvetica", 12)
            y = top - 50
            continue
        text = info.paragraph.text
        if idx == skip_empty:
            y -= 24
            continue
        c.setFont("Helvetica", 1 if text.startswith(dl.MARK) else 12)
        c.drawString(72, y, text)
        c.setFont("Helvetica", 12)
        y -= 24
    c.save()


class DocxLayoutTests(unittest.TestCase):
    def setUp(self):
        import docx
        from docx_ai import docx_layout as dl
        self.dl = dl
        self.doc = _sample_doc()
        self.copy = docx.Document(io.BytesIO(dl.render_copy(self.doc)))
        self.pdf = Path(tempfile.mkdtemp()) / "x.pdf"

    def test_copy_marks_only_empty_paragraphs_uniquely(self):
        texts = [i.paragraph.text for i in self.dl.paragraphs(self.copy)]
        originals = [i.paragraph.text for i in self.dl.paragraphs(self.doc)]
        self.assertEqual(len(texts), len(originals))
        marks = [t for t, o in zip(texts, originals) if not o]
        self.assertEqual(len(marks), len(set(marks)))
        self.assertTrue(all(m.startswith(self.dl.MARK) for m in marks))
        self.assertEqual([t for t, o in zip(texts, originals) if o], [o for o in originals if o])

    def test_marker_keeps_line_height(self):
        """Il segno ha la stessa dimensione del segno di paragrafo (altezza di riga invariata),
        e' bianco e largo l'1%; elementi di w:rPr nell'ordine dello schema."""
        import docx
        from docx.oxml.ns import qn
        from docx.shared import Pt
        d = docx.Document()
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        rpr = p._p.get_or_add_pPr().get_or_add_rPr() if hasattr(p._p.get_or_add_pPr(), "get_or_add_rPr") \
            else None
        if rpr is None:
            from docx.oxml import OxmlElement
            rpr = OxmlElement("w:rPr")
            p._p.get_or_add_pPr().append(rpr)
        sz = rpr.makeelement(qn("w:sz"), {qn("w:val"): "36"})
        rpr.append(sz)
        copy = docx.Document(io.BytesIO(self.dl.render_copy(d)))
        r = copy.paragraphs[0].runs[0]._r
        names = [c.tag.split("}")[1] for c in r.find(qn("w:rPr"))]
        self.assertEqual(names, ["color", "w", "sz"])
        self.assertEqual(r.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")), "36")

    def test_every_paragraph_located(self):
        _fake_render(self.copy, self.pdf)
        infos = self.dl.paragraphs(self.doc)
        lay = self.dl.build_layout(self.pdf, self.dl.snapshot(infos))
        self.assertEqual(sorted(lay.paras), list(range(len(infos))))
        # cella vuota: rettangolo della cella dai bordi della tabella
        empty_cell = next(i for i, x in enumerate(infos) if x.in_cell and not x.paragraph.text)
        x0, y0, x1, y1 = lay.paras[empty_cell].area
        self.assertAlmostEqual(x1 - x0, 198, delta=3)
        self.assertEqual(self.dl.locate(lay, 0, (x0 + x1) / 2, (y0 + y1) / 2), (empty_cell, 0))
        # riga da compilare "______" dopo "Cliente:"
        b = lay.paras[1].boxes[12]
        i, off = self.dl.locate(lay, 0, (b[0] + b[2]) / 2, (b[1] + b[3]) / 2)
        self.assertEqual(i, 1)
        self.assertIn(off, (12, 13))
        self.assertEqual(len(self.dl.span_boxes(lay.paras[1], 0, 8)), 1)  # "Cliente:" su una riga

    def test_missing_empty_paragraph_does_not_derail(self):
        infos = self.dl.paragraphs(self.doc)
        _fake_render(self.copy, self.pdf, skip_empty=2)
        lay = self.dl.build_layout(self.pdf, self.dl.snapshot(infos))
        self.assertNotIn(2, lay.paras)
        self.assertEqual(sorted(lay.paras), [i for i in range(len(infos)) if i != 2])


class XlsxLayoutTests(unittest.TestCase):
    def test_geometry_and_styles(self):
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from docx_ai import xlsx_layout as xl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["C"].hidden = True
        ws.row_dimensions[2].height = 30
        ws["A1"] = "Cliente"
        ws["A1"].font = Font(bold=True, color="FF0000", size=14)
        ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
        ws["B2"] = 1250.5
        ws["B2"].number_format = "#,##0.00"
        ws["A3"].border = Border(left=Side("thin"), bottom=Side("double", color="FF0000"))
        ws.merge_cells("A4:B5")
        ws["A4"] = "Note"
        ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
        buf = io.BytesIO()
        wb.save(buf)
        wb2 = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        m = xl.build(wb2.active, xl.theme_colors(wb2))
        self.assertEqual(m.xs[1], 140)          # 20 caratteri
        self.assertEqual(m.xs[2] - m.xs[1], 64)  # colonna predefinita: 64 px come Excel
        self.assertEqual(m.xs[3], m.xs[2])       # colonna C nascosta
        self.assertEqual(m.ys[2] - m.ys[1], 40)  # 30 pt
        a1 = m.cells[(1, 1)]
        self.assertEqual((a1.font[2], a1.color), (True, "#FF0000"))
        self.assertEqual(m.cells[(1, 2)].fill, "#FFFF00")
        self.assertEqual(m.cells[(2, 2)].text, "1.250,50")
        self.assertEqual(m.cells[(2, 2)].halign, "right")
        self.assertEqual(m.cells[(3, 1)].borders["bottom"][1], "#FF0000")
        self.assertEqual(xl.cell_box(m, (4, 1)), (0, m.ys[3], m.xs[2], m.ys[5]))
        self.assertEqual(xl.cell_at(m, 150, m.ys[4] + 2), (4, 1))  # cella unita -> alto a sinistra
        self.assertGreater(a1.clip_right, a1.box[2])               # testo che sborda a destra

    def test_theme_tint(self):
        from openpyxl.styles.colors import Color
        from docx_ai import xlsx_layout as xl
        self.assertEqual(xl.resolve_color(Color(theme=1), xl.DEFAULT_THEME), "#000000")
        light = xl.resolve_color(Color(theme=4, tint=0.8), xl.DEFAULT_THEME)
        self.assertNotEqual(light, "#4472C4")


if __name__ == "__main__":
    unittest.main()


class FaithfulPdfTests(unittest.TestCase):
    def test_pdf_from_module_with_notes_appended(self):
        from unittest import mock
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader
        from docx_ai.services import report_service as rs

        class FakeConv:
            def can_convert(self, suffix):
                return True

            def convert(self, src, dst, timeout=120.0):
                c = canvas.Canvas(str(dst))
                c.drawString(72, 700, "MODULO ORIGINALE")
                c.save()

        tmp = Path(tempfile.mkdtemp())
        doc, pdf = tmp / "m.docx", tmp / "m.pdf"
        doc.write_bytes(b"x")
        mod = mock.Mock(schema={"type": "object", "properties": {"a": {"type": "string"}}}, name="M")
        with mock.patch("docx_ai.docx_layout.converter", return_value=FakeConv()), \
                mock.patch.object(rs, "_export_pdf_isolated",
                                  side_effect=lambda path, *a, **k: FakeConv().convert(None, path)):
            ok = rs._faithful_pdf(doc, pdf, {"a": "x"}, mod, 1, "nota", [])
        self.assertTrue(ok)
        pages = PdfReader(str(pdf)).pages
        self.assertEqual(len(pages), 2)
        self.assertIn("MODULO ORIGINALE", pages[0].extract_text())

    def test_without_engine_falls_back(self):
        from unittest import mock
        from docx_ai.services import report_service as rs
        conv = mock.Mock()
        conv.can_convert.return_value = False
        with mock.patch("docx_ai.docx_layout.converter", return_value=conv):
            self.assertFalse(rs._faithful_pdf(Path("x.docx"), Path("x.pdf"), {}, mock.Mock(), 1, "", []))
