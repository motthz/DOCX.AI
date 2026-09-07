"""Regression tests for bugs found during MASTER AUTONOMOUS CODEBASE AUDIT.

Each test corresponds to an AUDIT-XXX entry in AUDIT_PROGRESS.md.
The tests are designed to FAIL before the fix and PASS after.
"""

from __future__ import annotations

import shutil
import sqlite3
import sys
import tempfile
import threading
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from maintenance_ai.db import Database  # noqa: E402
from maintenance_ai.config import Config  # noqa: E402
from maintenance_ai.security import SecurityError  # noqa: E402
from maintenance_ai.module_manager import ModuleManager  # noqa: E402
from maintenance_ai.llm.llama_server import LlamaServer, LlamaServerOptions  # noqa: E402


class TestAudit010SqliteThreadSafety(unittest.TestCase):
    """AUDIT-010 P1 CRITICAL: SQLite ProgrammingError under multi-threaded access."""

    def setUp(self) -> None:
        self._tmp = Path(tempfile.mkdtemp())
        self.db_path = self._tmp / "audit010.sqlite3"
        self.addCleanup(lambda: shutil.rmtree(self._tmp, ignore_errors=True))

    def test_concurrent_read_write_no_programming_error(self) -> None:
        db = Database(self.db_path)
        errors: list = []
        lock = threading.Lock()

        def worker(tid: int) -> None:
            try:
                for i in range(15):
                    db.set_setting(f"k{tid}_{i}", f"v{tid}_{i}")
                    got = db.get_setting(f"k{tid}_{i}")
                    self.assertEqual(got, f"v{tid}_{i}")
                    db.upsert_module(
                        slug=f"m{tid}_{i}",
                        name=f"N{tid}_{i}",
                        template_type="docx",
                        template_path="t",
                        schema_path="s",
                        mapping_path="m",
                        module_json_path="mj",
                        folder_path="f",
                    )
                    rep = db.create_report(
                        module_id=1,
                        module_version="1.0.0",
                        status="draft",
                        input_description=f"t{tid}_{i}",
                    )
                    db.update_report(rep, status="approved")
            except Exception as exc:  # pragma: no cover
                with lock:
                    errors.append((tid, type(exc).__name__, str(exc)[:200]))

        threads = [threading.Thread(target=worker, args=(i,), daemon=True) for i in range(4)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(timeout=30.0)
        db.close()
        if errors:  # pragma: no cover - printed for visibility
            for e in errors:
                print("THREAD-ERROR:", e)
        self.assertEqual(errors, [])


class TestAudit001SqlWhitelist(unittest.TestCase):
    """AUDIT-001 + AUDIT-007 P2 HIGH: Reject invalid column names in dynamic SQL."""

    def setUp(self) -> None:
        self._tmp = Path(tempfile.mkdtemp())
        self.db = Database(self._tmp / "audit001.sqlite3")
        self.addCleanup(lambda: shutil.rmtree(self._tmp, ignore_errors=True))

    def tearDown(self) -> None:
        self.db.close()

    def test_upsert_module_rejects_invalid_column(self) -> None:
        with self.assertRaises(ValueError):
            self.db.upsert_module(
                slug="ok",
                name="OK",
                template_type="docx",
                template_path="t",
                schema_path="s",
                mapping_path="m",
                module_json_path="mj",
                folder_path="f",
                **{"name); DROP TABLE modules; --": "x"},
            )

    def test_upsert_document_rejects_invalid_column(self) -> None:
        with self.assertRaises(ValueError):
            self.db.upsert_document(
                module_id=1,
                kind="template",
                path="p",
                sha256="0" * 64,
                size_bytes=0,
                **{"bogus_col": 1},
            )

    def test_create_report_rejects_invalid_column(self) -> None:
        with self.assertRaises(ValueError):
            self.db.create_report(
                module_id=1,
                module_version="1.0.0",
                status="draft",
                input_description="test",
                **{"bad`col": "nope"},
            )

    def test_update_report_rejects_invalid_column(self) -> None:
        with self.assertRaises(ValueError):
            self.db.update_report(1, **{"\"; --": "x"})

    def test_valid_keys_accepted(self) -> None:
        mid = self.db.upsert_module(
            slug="good",
            name="G",
            template_type="docx",
            template_path="t",
            schema_path="s",
            mapping_path="m",
            module_json_path="mj",
            folder_path="f",
        )
        self.assertGreater(mid, 0)


class TestAudit006NestedTransaction(unittest.TestCase):
    """AUDIT-006 P2 HIGH: Nested transaction must not commit outer work prematurely."""

    def setUp(self) -> None:
        self._tmp = Path(tempfile.mkdtemp())
        self.db = Database(self._tmp / "audit006.sqlite3")
        self.addCleanup(lambda: shutil.rmtree(self._tmp, ignore_errors=True))

    def tearDown(self) -> None:
        self.db.close()

    def test_inner_exception_rolls_only_inner_via_savepoint(self) -> None:
        with self.db.transaction() as conn:
            self.db.set_setting("k1", "v1")
            try:
                with self.db.transaction() as _:
                    conn.execute(
                        "INSERT INTO settings(key,value) VALUES(?,?)",
                        ("inner_k", "inner_v"),
                    )
                    raise RuntimeError("inner boom")
            except RuntimeError:
                pass
            conn.execute(
                "INSERT INTO settings(key,value) VALUES(?,?)", ("after_inner", "ok")
            )

        self.assertEqual(self.db.get_setting("k1"), "v1")
        self.assertIsNone(self.db.get_setting("inner_k"))
        self.assertEqual(self.db.get_setting("after_inner"), "ok")


class TestAudit002And009ModuleCreateAtomic(unittest.TestCase):
    """AUDIT-002 + AUDIT-009 P3 MEDIUM: mkdir TOCTOU race raises SecurityError not FileExistsError."""

    def setUp(self) -> None:
        import os
        self._tmp = Path(tempfile.mkdtemp())
        self._prev_data = os.environ.get("MAINTENANCE_AI_DATA_DIR")
        self._prev_app = os.environ.get("MAINTENANCE_AI_APP_DIR")
        data_dir = self._tmp / "data"
        data_dir.mkdir(parents=True, exist_ok=True)
        os.environ["MAINTENANCE_AI_DATA_DIR"] = str(data_dir)
        os.environ["MAINTENANCE_AI_APP_DIR"] = str(
            Path(__file__).resolve().parent.parent
        )
        self.addCleanup(self._cleanup_env)
        self.cfg = Config.load()
        self.db = Database(self.cfg.db_path())
        self.mm = ModuleManager(self.cfg, self.db)
        self.addCleanup(lambda: shutil.rmtree(self._tmp, ignore_errors=True))

    def _cleanup_env(self) -> None:
        import os
        for key, prev in (
            ("MAINTENANCE_AI_DATA_DIR", self._prev_data),
            ("MAINTENANCE_AI_APP_DIR", self._prev_app),
        ):
            if prev is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = prev

    def tearDown(self) -> None:
        self.db.close()

    def test_create_module_duplicate_slug_raises_security_error(self) -> None:
        self.mm.create_module("Modulo Primo", "docx", slug="slug1")
        with self.assertRaises(SecurityError):
            self.mm.create_module("Modulo Secondo", "docx", slug="slug1")
        # Now pre-create slug folder on disk (simulate concurrent mkdir winning the race)
        pre = self.cfg.workspace_root() / "precreated"
        pre.mkdir(parents=True, exist_ok=False)
        with self.assertRaises(SecurityError):
            self.mm.create_module("Race", "docx", slug="precreated")

    def test_create_module_from_template_duplicate_raises_security_error(self) -> None:
        self.mm.create_module("First", "docx", slug="dup-tpl")
        import docx
        from docx.shared import Pt
        tpl = self._tmp / "t.docx"
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("{{prova_field}}")
        d.save(str(tpl))
        with self.assertRaises(SecurityError):
            self.mm.create_module_from_template("Second", tpl, slug="dup-tpl")


class TestAudit008LlamaServerStdoutDrain(unittest.TestCase):
    """AUDIT-008 P2 HIGH: LlamaServer must start a drain thread when capture_logs is True."""

    def test_reader_thread_attribute_exists(self) -> None:
        opts = LlamaServerOptions(
            runtime_dir=Path("/nonexistent"),
            model_path=Path("/nonexistent/model.gguf"),
            capture_logs=True,
        )
        srv = LlamaServer(opts)
        self.assertTrue(hasattr(srv, "_stdout_drain_thread"))
        self.assertTrue(hasattr(srv, "_stdout_buffer"))
        self.assertIsNone(srv._stdout_drain_thread)
        self.assertEqual(srv._stdout_buffer, b"")


if __name__ == "__main__":  # pragma: no cover
    unittest.main(verbosity=2)
