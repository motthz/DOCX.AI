"""Logica dell'editor visuale: inserimento dei campi nei paragrafi DOCX."""

from __future__ import annotations

import unittest

import docx

from docx_ai import template_tools as tt


def _par(*runs: str):
    p = docx.Document().add_paragraph()
    for r in runs:
        p.add_run(r)
    return p


class FieldKeyTests(unittest.TestCase):
    def test_key_from_label(self):
        self.assertEqual(tt.field_key("Nome e cognome del cliente"), "nome_e_cognome_del_cliente")
        self.assertEqual(tt.field_key("N° ordine"), "n_ordine")
        self.assertEqual(tt.field_key("Città / Località"), "citta_localita")
        self.assertEqual(tt.field_key("2° turno"), "campo_2_turno")
        self.assertEqual(tt.field_key("!!!"), "campo")

    def test_key_is_unique(self):
        self.assertEqual(tt.field_key("Data", ["data", "data_2"]), "data_3")

    def test_suggest_label(self):
        self.assertEqual(tt.suggest_label("Data intervento: "), "Data intervento")
        self.assertEqual(tt.suggest_label("Cliente: {{cliente}}. Indirizzo "), "Indirizzo")
        self.assertEqual(tt.suggest_label(""), "")


class PlaceFieldTests(unittest.TestCase):
    def test_replaces_blank_line(self):
        p = _par("Nome: ", "__________", " fine")
        tt.place_field(p, 8, "nome")
        self.assertEqual(p.text, "Nome: {{nome}} fine")

    def test_inserts_with_spaces_and_keeps_words_whole(self):
        p = _par("Firma del tecnico")
        tt.place_field(p, 3, "firma")  # dentro "Firma": va alla fine della parola
        self.assertEqual(p.text, "Firma {{firma}} del tecnico")

    def test_replaces_selection_across_runs(self):
        p = _par("Cliente: Mar", "io Ros", "si srl")
        tt.place_field(p, 9, "cliente", end=24)
        self.assertEqual(p.text, "Cliente: {{cliente}}")

    def test_replaces_existing_placeholder(self):
        p = _par("Data: {{data}}")
        tt.place_field(p, 8, "data_visita")
        self.assertEqual(p.text, "Data: {{data_visita}}")

    def test_empty_paragraph(self):
        p = _par()
        tt.place_field(p, 0, "note")
        self.assertEqual(p.text, "{{note}}")

    def test_remove_and_rename(self):
        doc = docx.Document()
        a = doc.add_paragraph("A {{x}} B")
        b = doc.add_paragraph("{{x}} e {{y}}")
        self.assertEqual(tt.rename_field(doc.paragraphs, "x", "z"), 2)
        self.assertEqual((a.text, b.text), ("A {{z}} B", "{{z}} e {{y}}"))
        self.assertEqual(tt.remove_field(doc.paragraphs, "z"), 2)
        self.assertEqual((a.text, b.text), ("A B", " e {{y}}"))


class SpecTests(unittest.TestCase):
    def test_kind_round_trip(self):
        for kind in tt.FIELD_TYPES:
            self.assertEqual(tt.kind_of(tt.spec_for(kind, options=["A", "B"])), kind)

    def test_change_kind_keeps_title_and_description(self):
        spec = {"type": "string", "title": "Esito", "description": "Esito finale"}
        new = tt.change_kind(spec, "choice", ["Positivo", "Negativo"])
        self.assertEqual(new["enum"], ["Positivo", "Negativo"])
        self.assertEqual((new["title"], new["description"]), ("Esito", "Esito finale"))


if __name__ == "__main__":
    unittest.main()
