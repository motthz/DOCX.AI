"""Regressioni dell'audit "da utente" (ottobre 2026): ogni test riproduce un difetto
trovato usando l'app e fallisce senza la correzione corrispondente."""

from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

import docx  # noqa: E402

from docx_ai.parsers.docx_parser import apply_placeholders, extract_text  # noqa: E402


def _template(path: Path) -> Path:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Prot. {{numero_protocollo}}"
    p = d.add_paragraph("Dipendente: ")
    p.add_run("{{nome_")
    p.add_run("dipendente}}")  # segnaposto spezzato su due run, come fa Word
    d.add_paragraph("Città: {{città}} - {{campo sbagliato}}")
    outer = d.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "{{annidato}}"
    d.save(str(path))
    return path


class DocxPlaceholderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: shutil.rmtree(self.tmp, ignore_errors=True))
        self.path = _template(self.tmp / "t.docx")

    def test_finds_header_accented_nested_in_document_order(self):
        ext = extract_text(docx.Document(str(self.path)))
        self.assertEqual(ext.placeholders, ["nome_dipendente", "città", "annidato", "numero_protocollo"])
        self.assertEqual(ext.meta["invalid_placeholders"], ["{{campo sbagliato}}"])

    def test_replaces_everywhere(self):
        d = docx.Document(str(self.path))
        apply_placeholders(d, {"numero_protocollo": "42", "nome_dipendente": "Rossi", "città": "Milano",
                               "annidato": "ok"})
        self.assertEqual(d.sections[0].header.paragraphs[0].text, "Prot. 42")
        body = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Dipendente: Rossi", body)
        self.assertIn("Città: Milano", body)
        self.assertEqual(d.tables[0].cell(0, 0).tables[0].cell(0, 0).text, "ok")

    def test_linked_header_is_not_created(self):
        d = docx.Document()
        d.add_paragraph("{{a}}")
        self.assertTrue(d.sections[0].first_page_header.is_linked_to_previous)
        extract_text(d)
        self.assertTrue(d.sections[0].first_page_header.is_linked_to_previous)


class PipelineTests(unittest.TestCase):
    def test_empty_model_answer_does_not_crash(self):
        from docx_ai.llm.json_pipeline import _lenient_json_parse
        self.assertEqual(_lenient_json_parse(""), (None, "Risposta vuota dal modello"))
        parsed, err = _lenient_json_parse("niente json")
        self.assertIsNone(parsed)
        self.assertTrue(err)

    def test_quality_full_form_without_enums_is_100(self):
        from docx_ai.llm.quality_scorer import score
        schema = {"type": "object", "properties": {"nome": {"type": "string"}, "data": {"type": "string"}},
                  "required": ["nome", "data"]}
        self.assertEqual(score({"nome": "Rossi", "data": "05/10/2026"}, schema), 100)
        self.assertEqual(score({"nome": "Rossi", "data": "NON_SPECIFICATO"}, schema), 50)

    def test_ollama_disables_thinking(self):
        from docx_ai.llm.ollama_backend import OllamaBackend, OllamaBackendOptions
        sent = {}

        class Resp:
            status = 200

            def read(self):
                return b'{"choices": [{"message": {"content": "{}"}}]}'

            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

        def fake_urlopen(req, timeout=0):
            sent.update(json.loads(req.data.decode("utf-8")))
            return Resp()

        with mock.patch("urllib.request.urlopen", fake_urlopen):
            OllamaBackend(OllamaBackendOptions(model="qwen3:4b")).chat_completions([{"role": "user", "content": "x"}])
        self.assertEqual(sent.get("reasoning_effort"), "none")


class AppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp(prefix="mai_ua_"))
        cls.env = mock.patch.dict(os.environ, {"DOCX_AI_DATA_DIR": str(cls.tmp)})
        cls.env.start()
        from docx_ai.app import App
        cls.app = App.bootstrap()

    @classmethod
    def tearDownClass(cls):
        cls.app.shutdown()
        cls.env.stop()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_install_examples(self):
        mm = self.app.module_manager
        mm.install_examples()
        self.assertTrue({"Verbale_di_riunione", "Richiesta_acquisto"} <= {m.slug for m in mm.list_modules()})
        self.assertEqual(mm.install_examples(), [])  # gia' presenti: nessuna copia doppia

    def test_module_from_template_keeps_order_and_reports_invalid(self):
        tpl = _template(self.tmp / "ferie.docx")
        mod = self.app.module_manager.create_module_from_template(name="Ferie", template_file=tpl)
        self.assertEqual(list(mod.schema["properties"]),
                         ["nome_dipendente", "città", "annidato", "numero_protocollo"])
        self.assertEqual(self.app.module_manager.last_invalid_placeholders, ["{{campo sbagliato}}"])

    def test_import_existing_module_does_not_nest(self):
        mm = self.app.module_manager
        mm.install_examples()
        z = self.tmp / "verbale.zip"
        mm.export_module_zip("Verbale_di_riunione", z)
        mod = mm.import_module_zip(z)
        self.assertNotEqual(mod.slug, "Verbale_di_riunione")
        inside = [p.name for p in (mm.workspace / "Verbale_di_riunione").iterdir()]
        self.assertFalse([n for n in inside if n.startswith("__import__")], inside)

    def test_exported_documents_hide_sentinel(self):
        from docx_ai.services.report_service import _for_document
        self.assertEqual(_for_document({"a": "NON_SPECIFICATO", "b": ["x", "NON_SPECIFICATO"], "c": 3}),
                         {"a": "", "b": ["x"], "c": 3})

    def test_structured_extraction_sends_request_and_sources(self):
        from docx_ai.llm.json_pipeline import JsonPipeline
        from docx_ai.llm.llama_server import MockLlamaServer
        seen = []

        def answer(messages, json_schema=None):
            seen.append(messages)
            return {"titolo": "Procedura ferie", "sezioni": []}

        ai = self.app.ai_service
        schema = {"type": "object", "properties": {"titolo": {"type": "string"},
                                                   "sezioni": {"type": "array", "items": {"type": "object"}}},
                  "required": ["titolo"]}
        with mock.patch.object(ai, "pipeline", lambda **k: JsonPipeline(MockLlamaServer(answer), max_retries=0)):
            res = ai.extract_structured("document_creation", schema, "Istruzione per chiedere le ferie")
        self.assertTrue(res.success)
        self.assertEqual(res.data["titolo"], "Procedura ferie")
        self.assertIn("Istruzione per chiedere le ferie", seen[0][-1]["content"])

    def test_archived_modules_listed(self):
        mm = self.app.module_manager
        mm.install_examples()
        mm.archive_module("Richiesta_acquisto")
        self.assertIn("Richiesta_acquisto", [f for f, _n in mm.list_archived()])
        mm.restore_module("Richiesta_acquisto")
        self.assertNotIn("Richiesta_acquisto", [f for f, _n in mm.list_archived()])


class DocumentFeatureTests(unittest.TestCase):
    def test_markdown_to_docx_keeps_subsections(self):
        from docx_ai.docintelligence.document_generator import _markdown_to_docx
        d = docx.Document()
        _markdown_to_docx(d, "# Titolo\n\n## Sezione\nTesto\n### Sotto\n- punto")
        self.assertEqual([p.text for p in d.paragraphs], ["Titolo", "Sezione", "Testo", "Sotto", "punto"])
        self.assertEqual(d.paragraphs[3].style.name, "Heading 3")

    def test_source_refs_serializable(self):
        from docx_ai.docintelligence.source_tracker import Conflict, SourceRef
        ref = SourceRef(file_name="a.pdf", page=1, excerpt="x")
        json.dumps(ref.to_dict())
        json.dumps(Conflict(field="f", alternatives=[("1", [ref])]).to_dict())

    def test_audit_string_instead_of_list(self):
        from docx_ai.docintelligence.audit_engine import _as_list
        self.assertEqual(_as_list("regolamento.txt"), ["regolamento.txt"])
        self.assertEqual(_as_list(["a", "NON_SPECIFICATO", ""]), ["a"])
        self.assertEqual(_as_list(None), [])


class SchemaEditorTests(unittest.TestCase):
    def test_save_without_changes_keeps_every_field(self):
        import tkinter as tk
        from docx_ai.ui.schema_editor_dialog import SchemaEditorDialog
        try:
            root = tk.Tk()
        except tk.TclError:
            self.skipTest("Tk non disponibile")
        root.withdraw()
        try:
            for mod in ("Rapporto_Manutenzione_Citterio", "Verbale_di_riunione", "Richiesta_acquisto"):
                schema = json.loads((REPO / "examples" / "modules" / mod / "schema.json").read_text(encoding="utf-8"))
                dlg = SchemaEditorDialog(root, schema)
                out = dlg._collect_schema()
                dlg.top.destroy()
                self.assertEqual(out["properties"], schema["properties"], mod)
                self.assertEqual(out["required"], schema["required"], mod)
        finally:
            root.destroy()


class ExamplesBundledTests(unittest.TestCase):
    def test_spec_bundles_examples(self):
        spec = (REPO / "packaging" / "DOCX.AI.spec").read_text(encoding="utf-8")
        self.assertIn('"examples"', spec)
        self.assertTrue((REPO / "examples" / "modules" / "Verbale_di_riunione" / "module.json").is_file())


if __name__ == "__main__":
    unittest.main()
