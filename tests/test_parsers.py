"""Test su parsers DOCX/XLSX e gestione placeholder spezzati tra run."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
SRC = REPO / "src"
sys.path.insert(0, str(SRC))
sys.path.insert(0, str(REPO))


class DocxParserTests(unittest.TestCase):
    def test_placeholder_even_split_across_runs(self):
        import docx
        from maintenance_ai.parsers.docx_parser import (
            apply_placeholders,
            extract_text,
        )
        d = docx.Document()
        p = d.add_paragraph()
        # Simula Word che spezza {{impianto}}
        p.add_run("{{im")
        p.add_run("pia")
        p.add_run("nto}}")
        outdir = Path(tempfile.mkdtemp(prefix="mai_docx_"))
        src = outdir / "t.docx"
        d.save(str(src))

        loaded = docx.Document(str(src))
        self.assertIn("impianto", extract_text(loaded).placeholders)
        apply_placeholders(loaded, {"impianto": "Pompa XX"})
        saved = outdir / "out.docx"
        loaded.save(str(saved))

        reloaded = docx.Document(str(saved))
        full = "\n".join(par.text for par in reloaded.paragraphs)
        self.assertIn("Pompa XX", full)
        self.assertNotIn("{{impianto}}", full)

    def test_table_and_placeholders(self):
        import docx
        from maintenance_ai.parsers.docx_parser import (
            apply_placeholders,
            extract_text,
        )
        d = docx.Document()
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "note"
        t.cell(0, 1).text = "{{note}}"
        t.cell(1, 0).text = "esito"
        t.cell(1, 1).text = "{{esito}}"
        ext = extract_text(d)
        self.assertEqual(set(ext.placeholders), {"note", "esito"})
        apply_placeholders(d, {"note": "ciao", "esito": "regolare"})
        outdir = Path(tempfile.mkdtemp(prefix="mai_docx2_"))
        p = outdir / "out.docx"
        d.save(str(p))
        self.assertTrue(p.stat().st_size > 0)


class XlsxParserTests(unittest.TestCase):
    def test_mapping_writes(self):
        from openpyxl import Workbook
        from maintenance_ai.parsers.xlsx_parser import apply_mapping, extract_text
        wb = Workbook()
        ws = wb.active
        ws.title = "Rapporto"
        ws["B4"] = ""
        ws["F4"] = ""
        ws["B11"] = ""
        mapping = {
            "impianto": {"type": "cell", "sheet": "Rapporto", "cell": "B4"},
            "data_intervento": {"type": "cell", "sheet": "Rapporto", "cell": "F4"},
            "attivita_eseguite": {"type": "joined_cell", "sheet": "Rapporto",
                                 "cell": "B11", "separator": "\n"},
        }
        values = {
            "impianto": "Pompa PP",
            "data_intervento": "2026-08-27",
            "attivita_eseguite": ["A", "B", "C"],
        }
        applied = apply_mapping(wb, mapping, values)
        self.assertIn("impianto", applied)
        self.assertEqual(ws["B4"].value, "Pompa PP")
        self.assertIn("A\nB\nC", str(ws["B11"].value))
        # extract_text must see placeholder if template had any
        ws2 = wb.create_sheet("Altro")
        ws2["A1"] = "{{note}}"
        ext = extract_text(wb)
        self.assertIn("note", ext.placeholders)


if __name__ == "__main__":
    unittest.main()
