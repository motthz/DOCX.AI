"""Test di exporter PDF, DOCX e XLSX con un ciclo roundtrip."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
SRC = REPO / "src"
sys.path.insert(0, str(SRC))
sys.path.insert(0, str(REPO))


SCHEMA = {
    "$schema": "https://json-schema.org/draft/2020-12/schema",
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "impianto": {"type": "string"},
        "esito": {"type": "string"},
        "attivita": {"type": "array", "items": {"type": "string"}},
        "note": {"type": "string"},
    },
    "required": ["impianto", "esito", "attivita", "note"],
}

VALUES = {
    "impianto": "Pompa X1",
    "esito": "regolare",
    "attivita": ["smontaggio", "sostituzione paraoli", "rimontaggio", "prova"],
    "note": "Tutto ok senza anomalie.",
}


class ExporterTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="mai_exp_"))

    def test_pdf_export(self):
        from maintenance_ai.exporters.pdf_exporter import export_pdf
        out = self.tmp / "out.pdf"
        export_pdf(out, VALUES, SCHEMA, report_id="1", module_name="Test")
        self.assertGreater(out.stat().st_size, 2000)
        with open(out, "rb") as f:
            head = f.read(5)
        self.assertEqual(head, b"%PDF-")

    def test_docx_roundtrip(self):
        import docx
        from maintenance_ai.exporters.docx_exporter import export_docx
        from maintenance_ai.parsers.docx_parser import extract_text
        # Crea template con placeholder
        tpl = self.tmp / "tpl.docx"
        d = docx.Document()
        d.add_paragraph("Impianto: {{impianto}}")
        d.add_paragraph("Esito: {{esito}}")
        d.add_paragraph("Attività:\n{{attivita}}")
        d.add_paragraph("Note: {{note}}")
        d.save(str(tpl))
        out = self.tmp / "out.docx"
        export_docx(tpl, out, VALUES)
        self.assertTrue(out.exists())
        ext = extract_text(docx.Document(str(out)))
        self.assertIn("Pompa X1", ext.full_text)
        self.assertIn("regolare", ext.full_text)
        self.assertNotIn("{{impianto}}", ext.full_text)

    def test_xlsx_roundtrip(self):
        from openpyxl import Workbook, load_workbook
        from maintenance_ai.exporters.xlsx_exporter import export_xlsx
        tpl = self.tmp / "tpl.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Rapporto"
        wb.save(str(tpl))
        mapping = {
            "impianto": {"type": "cell", "sheet": "Rapporto", "cell": "B2"},
            "esito": {"type": "cell", "sheet": "Rapporto", "cell": "B3"},
            "attivita": {"type": "joined_cell", "sheet": "Rapporto",
                         "cell": "B4", "separator": "\n"},
        }
        out = self.tmp / "out.xlsx"
        export_xlsx(tpl, out, mapping, VALUES)
        wb2 = load_workbook(str(out), data_only=True)
        ws2 = wb2["Rapporto"]
        self.assertEqual(str(ws2["B2"].value), "Pompa X1")
        self.assertIn("smontaggio", str(ws2["B4"].value))


if __name__ == "__main__":
    unittest.main()
